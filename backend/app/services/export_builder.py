"""
app/services/export_builder.py — Business logic for generating standardized CSV and XLSX exports.

Handles:
- Shared query building with soft-delete filtering
- Dynamic custom field discovery & deterministic column ordering
- Sanitization for spreadsheet formula injection in CSV
- Typed cell writing in XLSX (numbers, dates, percentages)
- Auto-fit column widths, freeze panes, autofilter in XLSX
- Exports for Invoices, Line Items, Receipts, and Unified Documents
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, literal, or_
from sqlalchemy.orm import Session, joinedload

from app.models.invoice import ImportBatch, InvoiceLineItem, InvoiceRecord, ReceiptLineItem, ReceiptRecord
from app.services.business_query import active_record_scope, invoice_filter_conditions, receipt_filter_conditions
from app.models.source_file import SourceFile
from app.models.template import Template


# ── Friendly Field Labels ─────────────────────────────────────────────────────

FRIENDLY_LABELS: dict[str, str] = {
    "company_name": "Company Name",
    "invoice_number": "Invoice Number",
    "invoice_date": "Invoice Date",
    "currency": "Currency",
    "total_amount": "Total Amount",
    "direction": "Direction",
    "merchant_name": "Merchant Name",
    "receipt_number": "Receipt Number",
    "transaction_date": "Transaction Date",
    "transaction_time": "Transaction Time",
    "payment_method": "Payment Method",
    "store_address": "Store Address",
    "description": "Description",
    "quantity": "Quantity",
    "unit_price": "Unit Price",
    "tax_rate": "Tax Rate",
    "tax_amount": "Tax Amount",
    "amount": "Amount",
    "due_date": "Due Date",
    "sub_total": "Subtotal",
    "subtotal": "Subtotal",
    "payment_terms": "Payment Terms",
    "discount": "Discount",
    "shipping_amount": "Shipping Amount",
    "shipping": "Shipping",
    "sku_code": "SKU Code",
    "sku": "SKU Code",
    "uom": "UOM",
    "unit_of_measure": "Unit of Measure",
    "po_number": "PO Number",
    "purchase_order_number": "PO Number",
    "source_file": "Source File",
    "template_name": "Template Name",
    "source_worksheet": "Worksheet",
    "source_row_number": "Source Row",
    "imported_at": "Imported At",
}


def friendly_label(key: str) -> str:
    """Returns a clean human-readable label for a canonical or custom key."""
    if key in FRIENDLY_LABELS:
        return FRIENDLY_LABELS[key]
    clean_key = str(key).strip().lower().replace("-", "_")
    if clean_key in FRIENDLY_LABELS:
        return FRIENDLY_LABELS[clean_key]
    return key.replace("_", " ").title()


# ── Formula Injection Protection ──────────────────────────────────────────────

DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_csv_value(val: Any) -> Any:
    """
    Guards against CSV spreadsheet formula injection for text values.
    If a text field begins with dangerous formula characters, prepends a single quote.
    Legitimate numbers and dates are untouched.
    """
    if val is None:
        return ""
    if isinstance(val, (int, float, Decimal)):
        return val
    if isinstance(val, (date, datetime)):
        return val.isoformat()
    if isinstance(val, bool):
        return "true" if val else "false"

    s_val = str(val)
    if s_val and s_val.startswith(DANGEROUS_FORMULA_PREFIXES):
        return f"'{s_val}"
    return s_val


# ── Shared Query Filter Builders ──────────────────────────────────────────────


def build_invoices_export_query(
    db: Session,
    search: str | None = None,
    company: str | None = None,
    currency: str | None = None,
    direction: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    has_line_items: bool | None = None,
):
    """
    Builds the SQLAlchemy query for invoice export with soft-delete filtering.
    """
    query = (
        db.query(InvoiceRecord)
        .options(
            joinedload(InvoiceRecord.batch).joinedload(ImportBatch.source_file),
            joinedload(InvoiceRecord.batch).joinedload(ImportBatch.template),
        )
        .join(InvoiceRecord.batch)
        .filter(active_record_scope())
    )

    if search and search.strip():
        search_term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                InvoiceRecord.company_name.ilike(search_term),
                InvoiceRecord.invoice_number.ilike(search_term),
            )
        )

    query = query.filter(*invoice_filter_conditions(company, currency, date_from, date_to, direction))

    if has_line_items is not None:
        if has_line_items:
            query = query.filter(InvoiceRecord.line_items.any())
        else:
            query = query.filter(~InvoiceRecord.line_items.any())

    return query.order_by(InvoiceRecord.invoice_date.desc().nulls_last(), InvoiceRecord.id.asc())


def build_receipts_export_query(
    db: Session,
    search: str | None = None,
    merchant: str | None = None,
    receipt_number: str | None = None,
    currency: str | None = None,
    direction: str | None = None,
    payment_method: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
):
    """
    Builds the SQLAlchemy query for receipt export with soft-delete filtering.
    """
    query = (
        db.query(ReceiptRecord)
        .options(
            joinedload(ReceiptRecord.batch).joinedload(ImportBatch.source_file),
            joinedload(ReceiptRecord.batch).joinedload(ImportBatch.template),
        )
        .join(ReceiptRecord.batch)
        .filter(active_record_scope())
    )

    if search and search.strip():
        search_term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                ReceiptRecord.merchant_name.ilike(search_term),
                ReceiptRecord.receipt_number.ilike(search_term),
            )
        )

    if receipt_number and receipt_number.strip():
        query = query.filter(ReceiptRecord.receipt_number.ilike(f"%{receipt_number.strip()}%"))

    query = query.filter(
        *receipt_filter_conditions(
            merchant=merchant,
            currency=currency,
            direction=direction,
            payment_method=payment_method,
            date_from=date_from,
            date_to=date_to,
        )
    )

    return query.order_by(ReceiptRecord.transaction_date.desc().nulls_last(), ReceiptRecord.id.asc())


def build_line_items_export_query(
    db: Session,
    search: str | None = None,
    company: str | None = None,
    currency: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
):
    """
    Builds the SQLAlchemy query for line items export with soft-delete filtering.
    """
    query = (
        db.query(InvoiceLineItem)
        .options(
            joinedload(InvoiceLineItem.invoice_record).joinedload(InvoiceRecord.batch).joinedload(ImportBatch.source_file),
        )
        .join(InvoiceLineItem.invoice_record)
        .join(InvoiceRecord.batch)
        .filter(active_record_scope())
    )

    if search and search.strip():
        search_term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                InvoiceLineItem.description.ilike(search_term),
                InvoiceRecord.company_name.ilike(search_term),
                InvoiceRecord.invoice_number.ilike(search_term),
            )
        )

    query = query.filter(*invoice_filter_conditions(company, currency, date_from, date_to))

    return query.order_by(
        InvoiceRecord.invoice_date.desc().nulls_last(),
        InvoiceRecord.id.asc(),
        InvoiceLineItem.source_row_number.asc(),
        InvoiceLineItem.id.asc(),
    )


def fetch_documents_export_data(
    db: Session,
    search: str | None = None,
    document_type: str | None = None,
    direction: str | None = None,
    currency: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[dict[str, Any]]:
    """
    Fetches unified financial document rows (invoices and receipts) for export.
    """
    inv_li_count = (
        db.query(InvoiceLineItem.invoice_id, func.count(InvoiceLineItem.id).label("cnt"))
        .group_by(InvoiceLineItem.invoice_id)
        .subquery()
    )
    rec_li_count = (
        db.query(ReceiptLineItem.receipt_id, func.count(ReceiptLineItem.id).label("cnt"))
        .group_by(ReceiptLineItem.receipt_id)
        .subquery()
    )

    queries = []
    include_invoices = document_type in {None, "", "all", "invoice"}
    include_receipts = document_type in {None, "", "all", "receipt"}

    if include_invoices:
        q_inv = (
            db.query(
                literal("Invoice").label("doc_type"),
                InvoiceRecord.id.label("record_id"),
                InvoiceRecord.invoice_number.label("document_number"),
                InvoiceRecord.company_name.label("counterparty"),
                InvoiceRecord.invoice_date.label("document_date"),
                InvoiceRecord.currency.label("currency"),
                InvoiceRecord.total_amount.label("total_amount"),
                InvoiceRecord.direction.label("direction"),
                SourceFile.original_name.label("source_filename"),
                ImportBatch.status.label("status"),
                ImportBatch.imported_at.label("imported_at"),
            )
            .join(InvoiceRecord.batch)
            .outerjoin(ImportBatch.source_file)
            .outerjoin(inv_li_count, InvoiceRecord.id == inv_li_count.c.invoice_id)
            .filter(active_record_scope())
        )
        queries.append(q_inv)

    if include_receipts:
        q_rec = (
            db.query(
                literal("Receipt").label("doc_type"),
                ReceiptRecord.id.label("record_id"),
                ReceiptRecord.receipt_number.label("document_number"),
                ReceiptRecord.merchant_name.label("counterparty"),
                ReceiptRecord.transaction_date.label("document_date"),
                ReceiptRecord.currency.label("currency"),
                ReceiptRecord.total_amount.label("total_amount"),
                ReceiptRecord.direction.label("direction"),
                SourceFile.original_name.label("source_filename"),
                ImportBatch.status.label("status"),
                ImportBatch.imported_at.label("imported_at"),
            )
            .join(ReceiptRecord.batch)
            .outerjoin(ImportBatch.source_file)
            .outerjoin(rec_li_count, ReceiptRecord.id == rec_li_count.c.receipt_id)
            .filter(active_record_scope())
        )
        queries.append(q_rec)

    if not queries:
        return []

    if len(queries) == 1:
        union_subq = queries[0].subquery()
    else:
        union_subq = queries[0].union_all(queries[1]).subquery()

    doc_query = db.query(union_subq)

    if search and search.strip():
        search_term = f"%{search.strip()}%"
        doc_query = doc_query.filter(
            or_(
                union_subq.c.counterparty.ilike(search_term),
                union_subq.c.document_number.ilike(search_term),
            )
        )

    if direction and direction.strip() and direction.strip().lower() != "all":
        doc_query = doc_query.filter(union_subq.c.direction == direction.strip().lower())

    if currency and currency.strip():
        doc_query = doc_query.filter(union_subq.c.currency == currency.strip().upper())

    if date_from:
        doc_query = doc_query.filter(union_subq.c.document_date >= date_from)

    if date_to:
        doc_query = doc_query.filter(union_subq.c.document_date <= date_to)

    doc_query = doc_query.order_by(union_subq.c.document_date.desc().nulls_last(), union_subq.c.record_id.desc())

    rows = doc_query.all()
    results = []
    for r in rows:
        results.append({
            "doc_type": r.doc_type,
            "record_id": r.record_id,
            "document_number": r.document_number or ("Not provided" if r.doc_type == "Receipt" else ""),
            "counterparty": r.counterparty,
            "document_date": r.document_date,
            "currency": r.currency,
            "total_amount": r.total_amount,
            "direction": r.direction or "unclassified",
            "source_filename": r.source_filename or "",
            "status": r.status or "imported",
            "imported_at": r.imported_at,
        })
    return results


# ── Custom Field Discovery ────────────────────────────────────────────────────

KNOWN_CUSTOM_ORDER = [
    "due_date",
    "sub_total",
    "subtotal",
    "tax_amount",
    "tax",
    "payment_terms",
    "discount",
    "shipping_amount",
    "shipping",
    "sku_code",
    "sku",
    "uom",
    "po_number",
]


def discover_custom_keys(records_custom_fields: list[dict[str, Any] | None]) -> list[str]:
    """
    Scans custom_fields across records and returns a deterministic, ordered list of keys.
    """
    discovered: set[str] = set()
    for cf in records_custom_fields:
        if isinstance(cf, dict):
            for k in cf.keys():
                if k:
                    discovered.add(k)

    def sort_key(k: str) -> tuple[int, str]:
        lower_k = k.lower().replace("-", "_")
        if lower_k in KNOWN_CUSTOM_ORDER:
            return (0, str(KNOWN_CUSTOM_ORDER.index(lower_k)))
        return (1, k.lower())

    return sorted(list(discovered), key=sort_key)


# ── CSV Exporters ─────────────────────────────────────────────────────────────


def generate_invoices_csv(records: list[InvoiceRecord]) -> str:
    """Generates standard CSV for invoice records."""
    custom_keys = discover_custom_keys([r.custom_fields for r in records])

    headers = [
        "Company Name",
        "Invoice Number",
        "Invoice Date",
        "Currency",
        "Total Amount",
        "Direction",
    ]
    for ck in custom_keys:
        headers.append(friendly_label(ck))
    headers.extend(["Source File", "Template Name", "Worksheet", "Imported At"])

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([sanitize_csv_value(v) for v in headers])

    for rec in records:
        batch = rec.batch
        src_file_name = batch.source_file.original_name if (batch and batch.source_file) else ""
        tmpl_name = batch.template.name if (batch and batch.template) else ""
        imported_at_str = batch.imported_at.strftime("%Y-%m-%d %H:%M:%S") if (batch and batch.imported_at) else ""

        row: list[Any] = [
            sanitize_csv_value(rec.company_name),
            sanitize_csv_value(rec.invoice_number),
            rec.invoice_date.isoformat() if rec.invoice_date else "",
            sanitize_csv_value(rec.currency or ""),
            str(rec.total_amount) if rec.total_amount is not None else "",
            sanitize_csv_value(rec.direction or "unclassified"),
        ]

        cf = rec.custom_fields or {}
        for ck in custom_keys:
            raw_val = cf.get(ck)
            if raw_val is None:
                row.append("")
            else:
                row.append(sanitize_csv_value(raw_val))

        row.extend([
            sanitize_csv_value(src_file_name),
            sanitize_csv_value(tmpl_name),
            sanitize_csv_value(rec.source_worksheet),
            imported_at_str,
        ])
        writer.writerow(row)

    return output.getvalue()


def generate_receipts_csv(records: list[ReceiptRecord]) -> str:
    """Generates standard CSV for receipt records."""
    headers = [
        "Merchant Name",
        "Receipt Number",
        "Transaction Date",
        "Transaction Time",
        "Direction",
        "Currency",
        "Subtotal",
        "Tax Amount",
        "Total Amount",
        "Payment Method",
        "Store Address",
        "Source File",
        "Imported At",
    ]

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([sanitize_csv_value(v) for v in headers])

    for rec in records:
        batch = rec.batch
        src_file_name = batch.source_file.original_name if (batch and batch.source_file) else ""
        imported_at_str = batch.imported_at.strftime("%Y-%m-%d %H:%M:%S") if (batch and batch.imported_at) else ""

        row: list[Any] = [
            sanitize_csv_value(rec.merchant_name),
            sanitize_csv_value(rec.receipt_number or "Not provided"),
            rec.transaction_date.isoformat() if rec.transaction_date else "",
            sanitize_csv_value(rec.transaction_time or ""),
            sanitize_csv_value(rec.direction or "unclassified"),
            sanitize_csv_value(rec.currency or ""),
            str(rec.subtotal) if rec.subtotal is not None else "",
            str(rec.tax_amount) if rec.tax_amount is not None else "",
            str(rec.total_amount) if rec.total_amount is not None else "",
            sanitize_csv_value(rec.payment_method or ""),
            sanitize_csv_value(rec.store_address or ""),
            sanitize_csv_value(src_file_name),
            imported_at_str,
        ]
        writer.writerow(row)

    return output.getvalue()


def generate_documents_csv(docs: list[dict[str, Any]]) -> str:
    """Generates standard CSV for unified document headers."""
    headers = [
        "Date",
        "Document Type",
        "Document Number",
        "Counterparty",
        "Direction",
        "Currency",
        "Total Amount",
        "Status",
        "Source File",
        "Imported At",
    ]

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([sanitize_csv_value(v) for v in headers])

    for d in docs:
        doc_date = d["document_date"].isoformat() if d["document_date"] else ""
        imported_at_str = d["imported_at"].strftime("%Y-%m-%d %H:%M:%S") if d.get("imported_at") else ""
        row = [
            doc_date,
            sanitize_csv_value(d["doc_type"]),
            sanitize_csv_value(d["document_number"]),
            sanitize_csv_value(d["counterparty"]),
            sanitize_csv_value(d["direction"]),
            sanitize_csv_value(d["currency"] or ""),
            str(d["total_amount"]) if d["total_amount"] is not None else "",
            sanitize_csv_value(d["status"]),
            sanitize_csv_value(d["source_filename"]),
            imported_at_str,
        ]
        writer.writerow(row)

    return output.getvalue()


def generate_line_items_csv(items: list[InvoiceLineItem]) -> str:
    """Generates standard CSV for flattened line items."""
    custom_keys = discover_custom_keys([li.custom_fields for li in items])

    headers = [
        "Company Name",
        "Invoice Number",
        "Invoice Date",
        "Currency",
        "Description",
        "Quantity",
        "Unit Price",
        "Tax Rate",
        "Tax Amount",
        "Amount",
    ]
    for ck in custom_keys:
        headers.append(friendly_label(ck))
    headers.extend(["Source File", "Source Row"])

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([sanitize_csv_value(v) for v in headers])

    for li in items:
        parent = li.invoice_record
        batch = parent.batch if parent else None
        src_file_name = batch.source_file.original_name if (batch and batch.source_file) else ""

        row: list[Any] = [
            sanitize_csv_value(parent.company_name if parent else ""),
            sanitize_csv_value(parent.invoice_number if parent else ""),
            parent.invoice_date.isoformat() if (parent and parent.invoice_date) else "",
            sanitize_csv_value((parent.currency if parent else "") or ""),
            sanitize_csv_value(li.description or ""),
            str(li.quantity) if li.quantity is not None else "",
            str(li.unit_price) if li.unit_price is not None else "",
            str(li.tax_rate) if li.tax_rate is not None else "",
            str(li.tax_amount) if li.tax_amount is not None else "",
            str(li.amount) if li.amount is not None else "",
        ]

        cf = li.custom_fields or {}
        for ck in custom_keys:
            raw_val = cf.get(ck)
            if raw_val is None:
                row.append("")
            else:
                row.append(sanitize_csv_value(raw_val))

        row.extend([
            sanitize_csv_value(src_file_name),
            li.source_row_number if li.source_row_number else "",
        ])
        writer.writerow(row)

    return output.getvalue()


# ── XLSX Exporters ────────────────────────────────────────────────────────────


def _style_xlsx_header(ws: openpyxl.worksheet.worksheet.Worksheet, headers: list[str]) -> None:
    """Applies clean Zen Data header styling, freeze panes, and autofilter."""
    header_font = Font(name="Arial", size=10, bold=True, color="1F2937")
    header_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    thin_border = Border(
        bottom=Side(style="medium", color="D1D5DB"),
        top=Side(style="thin", color="E5E7EB"),
    )

    ws.append(headers)
    for cell in ws[1]:
        cell.data_type = "s"
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"


def _auto_fit_columns(ws: openpyxl.worksheet.worksheet.Worksheet, min_width: int = 12, max_width: int = 45) -> None:
    """Adjusts column widths based on cell content length."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, max_width), min_width)


def _safe_float(val: Any) -> float | None:
    if val is None:
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def generate_invoices_xlsx(records: list[InvoiceRecord]) -> bytes:
    """Generates Excel workbook for invoice records with typed cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"

    custom_keys = discover_custom_keys([r.custom_fields for r in records])

    headers = [
        "Company Name",
        "Invoice Number",
        "Invoice Date",
        "Currency",
        "Total Amount",
        "Direction",
    ]
    for ck in custom_keys:
        headers.append(friendly_label(ck))
    headers.extend(["Source File", "Template Name", "Worksheet", "Imported At"])

    _style_xlsx_header(ws, headers)

    data_font = Font(name="Arial", size=10, color="111827")
    date_align = Alignment(horizontal="left")
    num_align = Alignment(horizontal="right")

    for rec in records:
        batch = rec.batch
        src_file_name = batch.source_file.original_name if (batch and batch.source_file) else None
        tmpl_name = batch.template.name if (batch and batch.template) else None
        imported_at_dt = batch.imported_at if (batch and batch.imported_at) else None

        row_cells: list[Any] = [
            rec.company_name,
            rec.invoice_number,
            rec.invoice_date,
            rec.currency,
            _safe_float(rec.total_amount),
            rec.direction or "unclassified",
        ]

        cf = rec.custom_fields or {}
        for ck in custom_keys:
            raw_val = cf.get(ck)
            if raw_val is None:
                row_cells.append(None)
            else:
                num_v = _safe_float(raw_val)
                if num_v is not None and not str(raw_val).startswith(("0", "+")) or str(raw_val) == "0":
                    row_cells.append(num_v)
                elif isinstance(raw_val, str) and len(raw_val) == 10 and raw_val[4] == "-" and raw_val[7] == "-":
                    try:
                        row_cells.append(date.fromisoformat(raw_val))
                    except ValueError:
                        row_cells.append(raw_val)
                else:
                    row_cells.append(raw_val)

        row_cells.extend([
            src_file_name,
            tmpl_name,
            rec.source_worksheet,
            imported_at_dt,
        ])

        ws.append(row_cells)
        row_idx = ws.max_row
        for col_idx, val in enumerate(row_cells, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(val, str):
                cell.data_type = "s"
            cell.font = data_font

            if isinstance(val, date) and not isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD"
                cell.alignment = date_align
            elif isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD HH:MM:SS"
                cell.alignment = date_align
            elif isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = num_align

    _auto_fit_columns(ws)

    buffer = io.BytesIO()
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def generate_receipts_xlsx(records: list[ReceiptRecord]) -> bytes:
    """Generates Excel workbook for receipt records with typed cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receipts"

    headers = [
        "Merchant Name",
        "Receipt Number",
        "Transaction Date",
        "Transaction Time",
        "Direction",
        "Currency",
        "Subtotal",
        "Tax Amount",
        "Total Amount",
        "Payment Method",
        "Store Address",
        "Source File",
        "Imported At",
    ]

    _style_xlsx_header(ws, headers)

    data_font = Font(name="Arial", size=10, color="111827")
    date_align = Alignment(horizontal="left")
    num_align = Alignment(horizontal="right")

    for rec in records:
        batch = rec.batch
        src_file_name = batch.source_file.original_name if (batch and batch.source_file) else None
        imported_at_dt = batch.imported_at if (batch and batch.imported_at) else None

        row_cells = [
            rec.merchant_name,
            rec.receipt_number or "Not provided",
            rec.transaction_date,
            rec.transaction_time,
            rec.direction or "unclassified",
            rec.currency,
            _safe_float(rec.subtotal),
            _safe_float(rec.tax_amount),
            _safe_float(rec.total_amount),
            rec.payment_method,
            rec.store_address,
            src_file_name,
            imported_at_dt,
        ]

        ws.append(row_cells)
        row_idx = ws.max_row
        for col_idx, val in enumerate(row_cells, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(val, str):
                cell.data_type = "s"
            cell.font = data_font

            if isinstance(val, date) and not isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD"
                cell.alignment = date_align
            elif isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD HH:MM:SS"
                cell.alignment = date_align
            elif isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = num_align

    _auto_fit_columns(ws)

    buffer = io.BytesIO()
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def generate_documents_xlsx(docs: list[dict[str, Any]]) -> bytes:
    """Generates Excel workbook for unified document records with typed cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Documents"

    headers = [
        "Date",
        "Document Type",
        "Document Number",
        "Counterparty",
        "Direction",
        "Currency",
        "Total Amount",
        "Status",
        "Source File",
        "Imported At",
    ]

    _style_xlsx_header(ws, headers)

    data_font = Font(name="Arial", size=10, color="111827")
    date_align = Alignment(horizontal="left")
    num_align = Alignment(horizontal="right")

    for d in docs:
        row_cells = [
            d["document_date"],
            d["doc_type"],
            d["document_number"],
            d["counterparty"],
            d["direction"],
            d["currency"],
            _safe_float(d["total_amount"]),
            d["status"],
            d["source_filename"],
            d["imported_at"],
        ]

        ws.append(row_cells)
        row_idx = ws.max_row
        for col_idx, val in enumerate(row_cells, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(val, str):
                cell.data_type = "s"
            cell.font = data_font

            if isinstance(val, date) and not isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD"
                cell.alignment = date_align
            elif isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD HH:MM:SS"
                cell.alignment = date_align
            elif isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = num_align

    _auto_fit_columns(ws)

    buffer = io.BytesIO()
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def generate_line_items_xlsx(items: list[InvoiceLineItem]) -> bytes:
    """Generates Excel workbook for flattened line items with typed cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Line Items"

    custom_keys = discover_custom_keys([li.custom_fields for li in items])

    headers = [
        "Company Name",
        "Invoice Number",
        "Invoice Date",
        "Currency",
        "Description",
        "Quantity",
        "Unit Price",
        "Tax Rate",
        "Tax Amount",
        "Amount",
    ]
    for ck in custom_keys:
        headers.append(friendly_label(ck))
    headers.extend(["Source File", "Source Row"])

    _style_xlsx_header(ws, headers)

    data_font = Font(name="Arial", size=10, color="111827")
    date_align = Alignment(horizontal="left")
    num_align = Alignment(horizontal="right")

    for li in items:
        parent = li.invoice_record
        batch = parent.batch if parent else None
        src_file_name = batch.source_file.original_name if (batch and batch.source_file) else None

        row_cells: list[Any] = [
            parent.company_name if parent else None,
            parent.invoice_number if parent else None,
            parent.invoice_date if parent else None,
            parent.currency if parent else None,
            li.description,
            _safe_float(li.quantity),
            _safe_float(li.unit_price),
            _safe_float(li.tax_rate),
            _safe_float(li.tax_amount),
            _safe_float(li.amount),
        ]

        cf = li.custom_fields or {}
        for ck in custom_keys:
            raw_val = cf.get(ck)
            if raw_val is None:
                row_cells.append(None)
            else:
                num_v = _safe_float(raw_val)
                if num_v is not None and not str(raw_val).startswith(("0", "+")) or str(raw_val) == "0":
                    row_cells.append(num_v)
                else:
                    row_cells.append(raw_val)

        row_cells.extend([
            src_file_name,
            li.source_row_number,
        ])

        ws.append(row_cells)
        row_idx = ws.max_row
        for col_idx, val in enumerate(row_cells, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(val, str):
                cell.data_type = "s"
            cell.font = data_font

            if isinstance(val, date) and not isinstance(val, datetime):
                cell.number_format = "YYYY-MM-DD"
                cell.alignment = date_align
            elif col_idx == 8 and isinstance(val, (int, float)):
                cell.number_format = "0.00%"
                cell.alignment = num_align
            elif col_idx == 6 and isinstance(val, (int, float)):
                cell.number_format = "#,##0.####"
                cell.alignment = num_align
            elif isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = num_align

    _auto_fit_columns(ws)

    buffer = io.BytesIO()
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
