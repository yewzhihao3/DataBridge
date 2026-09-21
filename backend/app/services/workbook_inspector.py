"""
app/services/workbook_inspector.py — Safe Excel workbook inspection.

This service is responsible for:
1. Validating that a file is a genuine, non-oversized, uncorrupted Excel (.xlsx) file.
2. Generating secure storage paths (UUID-based) with strict path-traversal prevention.
3. Inspecting workbook metadata: worksheet names, visibility, row/col bounds.
4. Detecting formula cells and distinguishing between evaluated vs. uncalculated formulas
   by dual-loading (data_only=False for formula strings, data_only=True for cached values).
5. Providing worksheet preview grids for user inspection.

──────────────────────────────────────────────────────────────────────────────
Architecture note:
This service is framework-independent. It does not import FastAPI Request/Response
objects or database sessions. It operates on Path objects and returns pure Python
dataclasses.
──────────────────────────────────────────────────────────────────────────────
"""

from __future__ import annotations

import os
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


# ── Constants ─────────────────────────────────────────────────────────────────

# Excel .xlsx files are ZIP archives starting with the standard PK magic bytes
ZIP_MAGIC_BYTES: bytes = b"PK\x03\x04"
ALLOWED_EXTENSIONS: frozenset[str] = frozenset({".xlsx"})
DEFAULT_MAX_FILE_SIZE_MB: int = 10


# ── Custom Exceptions ─────────────────────────────────────────────────────────


class WorkbookInspectionError(Exception):
    """Base exception for all workbook inspection and validation failures."""


class InvalidFileFormatError(WorkbookInspectionError):
    """Raised when file extension or magic bytes do not match valid .xlsx."""


class FileSizeExceededError(WorkbookInspectionError):
    """Raised when a file exceeds the allowed size limit."""


class PathTraversalError(WorkbookInspectionError):
    """Raised when an unsafe path attempts to escape the allowed base directory."""


class CorruptWorkbookError(WorkbookInspectionError):
    """Raised when openpyxl or zipfile fails to parse a corrupted workbook."""


class WorksheetNotFoundError(WorkbookInspectionError):
    """Raised when a requested worksheet is not present in the workbook."""


# ── Data Classes ──────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class FormulaCellDetail:
    """
    Detailed information about a formula cell found in a worksheet.

    Attributes:
        sheet_name: The name of the worksheet containing the formula.
        coordinate: Cell coordinate (e.g. 'D6').
        formula: The raw formula string (e.g. '=SUM(D2:D5)').
        cached_value: The evaluated value if cached by Excel, or None.
        is_cached_value_available: True if a non-None cached value exists.
    """

    sheet_name: str
    coordinate: str
    formula: str
    cached_value: Any
    is_cached_value_available: bool


@dataclass(frozen=True)
class WorksheetInfo:
    """
    Metadata about an individual worksheet inside a workbook.

    Attributes:
        name: Name of the worksheet.
        is_hidden: True if sheet_state is 'hidden' or 'veryHidden'.
        row_count: Number of rows with data/dimensions.
        column_count: Number of columns with data/dimensions.
        has_formulas: True if one or more formula cells exist in this sheet.
        formula_count: Total count of formula cells detected in this sheet.
    """

    name: str
    is_hidden: bool
    row_count: int
    column_count: int
    has_formulas: bool
    formula_count: int


@dataclass
class WorkbookInspection:
    """
    Complete inspection summary of an Excel workbook.

    Attributes:
        worksheets: List of WorksheetInfo for every sheet (visible and hidden).
        active_sheet_name: The name of the worksheet active by default.
        has_formula_cells: True if any worksheet contains formula cells.
        formula_cells: List of FormulaCellDetail for detected formulas.
        warnings: Human-readable warning messages (e.g. uncalculated formulas).
    """

    worksheets: list[WorksheetInfo]
    active_sheet_name: str
    has_formula_cells: bool
    formula_cells: list[FormulaCellDetail] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# ── File Validation and Secure Storage ────────────────────────────────────────


def generate_safe_upload_path(
    upload_dir: Path | str, original_filename: str
) -> tuple[Path, str]:
    """
    Generate a secure UUID-based storage path and validate against path traversal.

    1. Validates the original filename extension (.xlsx).
    2. Strips directory components from user-supplied names.
    3. Generates a unique UUID filename (e.g., '3f2b...8a.xlsx').
    4. Ensures the target path resolves strictly inside `upload_dir`.

    Args:
        upload_dir: Base directory where uploaded files are stored.
        original_filename: The raw filename provided by the client.

    Returns:
        A tuple of (resolved_destination_path, safe_unique_filename).

    Raises:
        InvalidFileFormatError: If the extension is not .xlsx.
        PathTraversalError: If the resolved path escapes upload_dir.
    """
    clean_name = os.path.basename(original_filename.strip())
    ext = os.path.splitext(clean_name)[1].lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise InvalidFileFormatError(
            f"Unsupported file extension '{ext}'. Only .xlsx files are supported."
        )

    base_dir = Path(upload_dir).resolve()
    unique_name = f"{uuid.uuid4().hex}{ext}"
    destination = (base_dir / unique_name).resolve()

    # Path traversal check: verify destination is a child of base_dir
    try:
        destination.relative_to(base_dir)
    except ValueError as exc:
        raise PathTraversalError(
            f"Destination path '{destination}' escapes base directory '{base_dir}'."
        ) from exc

    return destination, unique_name


def validate_xlsx_file(
    file_path: Path | str,
    max_size_mb: int = DEFAULT_MAX_FILE_SIZE_MB,
) -> None:
    """
    Perform pre-inspection validation on a local file.

    1. Checks file existence and ensures it is a regular file.
    2. Checks file size against max_size_mb.
    3. Checks magic bytes for ZIP/XLSX header signature (PK\\x03\\x04).

    Args:
        file_path: Path to the target file.
        max_size_mb: Maximum allowed file size in megabytes.

    Raises:
        FileNotFoundError: If the file does not exist.
        FileSizeExceededError: If the file exceeds max_size_mb.
        InvalidFileFormatError: If magic bytes or extension are invalid.
    """
    path = Path(file_path).resolve()

    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"File not found or not a regular file: '{path}'")

    file_size_bytes = path.stat().st_size
    max_size_bytes = max_size_mb * 1024 * 1024

    if file_size_bytes > max_size_bytes:
        actual_mb = file_size_bytes / (1024 * 1024)
        raise FileSizeExceededError(
            f"File size ({actual_mb:.2f} MB) exceeds maximum allowed size ({max_size_mb} MB)."
        )

    # Validate extension
    if path.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise InvalidFileFormatError(
            f"Invalid extension '{path.suffix}'. Only .xlsx files are allowed."
        )

    # Validate magic bytes (first 4 bytes of ZIP format)
    with open(path, "rb") as f:
        header = f.read(4)
        if header != ZIP_MAGIC_BYTES:
            raise InvalidFileFormatError(
                "Invalid file signature. The file does not appear to be a valid .xlsx archive."
            )


# ── Workbook Inspection ───────────────────────────────────────────────────────


def inspect_workbook(
    file_path: Path | str,
    max_size_mb: int = DEFAULT_MAX_FILE_SIZE_MB,
) -> WorkbookInspection:
    """
    Inspect an Excel workbook safely and extract structural metadata.

    Loads the workbook twice:
    - First with data_only=False to inspect formula syntax and worksheet states.
    - Second with data_only=True to inspect cached values of formulas.

    Gracefully catches corrupted files and invalid zip structures.

    Args:
        file_path: Path to the .xlsx file.
        max_size_mb: Maximum allowed file size in MB.

    Returns:
        WorkbookInspection containing worksheets, active sheet, formula details, and warnings.

    Raises:
        FileNotFoundError: If the file does not exist.
        FileSizeExceededError: If the file is too large.
        InvalidFileFormatError: If magic bytes/extension are invalid.
        CorruptWorkbookError: If openpyxl cannot parse the workbook.
    """
    validate_xlsx_file(file_path, max_size_mb=max_size_mb)
    path = Path(file_path).resolve()

    # 1. Load workbook with data_only=False (to detect formulas & structure)
    try:
        wb_formula = openpyxl.load_workbook(
            path, data_only=False, read_only=False, keep_vba=False
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError, Exception) as exc:
        raise CorruptWorkbookError(
            f"Failed to parse Excel workbook '{path.name}': {exc}"
        ) from exc

    # 2. Load workbook with data_only=True (to read cached formula values)
    try:
        wb_values = openpyxl.load_workbook(
            path, data_only=True, read_only=False, keep_vba=False
        )
    except Exception as exc:
        wb_formula.close()
        raise CorruptWorkbookError(
            f"Failed to read evaluated values from workbook '{path.name}': {exc}"
        ) from exc

    worksheets_info: list[WorksheetInfo] = []
    all_formula_cells: list[FormulaCellDetail] = []
    warnings: list[str] = []

    try:
        active_sheet_name = (
            wb_formula.active.title if wb_formula.active is not None else wb_formula.sheetnames[0]
        )

        for sheet_name in wb_formula.sheetnames:
            ws_formula = wb_formula[sheet_name]
            ws_values = wb_values[sheet_name]

            # Hidden status: 'visible', 'hidden', 'veryHidden'
            is_hidden = ws_formula.sheet_state != "visible"

            sheet_formula_count = 0
            max_row = ws_formula.max_row or 0
            max_col = ws_formula.max_column or 0

            # Scan cells for formulas
            for row in ws_formula.iter_rows():
                for cell in row:
                    val = cell.value
                    if val is None:
                        continue

                    is_formula = (
                        cell.data_type == "f"
                        or (isinstance(val, str) and val.startswith("="))
                    )

                    if is_formula:
                        sheet_formula_count += 1
                        coord = cell.coordinate
                        cached_val = ws_values[coord].value
                        has_cached = cached_val is not None

                        all_formula_cells.append(
                            FormulaCellDetail(
                                sheet_name=sheet_name,
                                coordinate=coord,
                                formula=str(val),
                                cached_value=cached_val,
                                is_cached_value_available=has_cached,
                            )
                        )

            has_formulas = sheet_formula_count > 0
            worksheets_info.append(
                WorksheetInfo(
                    name=sheet_name,
                    is_hidden=is_hidden,
                    row_count=max_row,
                    column_count=max_col,
                    has_formulas=has_formulas,
                    formula_count=sheet_formula_count,
                )
            )

        # Generate helpful warnings about uncalculated formula cells
        uncalculated_formulas = [
            f for f in all_formula_cells if not f.is_cached_value_available
        ]
        if uncalculated_formulas:
            sample_coords = ", ".join(
                f"'{f.sheet_name}!{f.coordinate}'" for f in uncalculated_formulas[:3]
            )
            count_text = (
                f"{len(uncalculated_formulas)} formula cell(s)"
                if len(uncalculated_formulas) > 1
                else "1 formula cell"
            )
            warnings.append(
                f"{count_text} (e.g. {sample_coords}) have no cached calculated values. "
                "If this workbook was generated without Excel formula evaluation, "
                "open and save it in Excel to compute cached results."
            )

        return WorkbookInspection(
            worksheets=worksheets_info,
            active_sheet_name=active_sheet_name,
            has_formula_cells=len(all_formula_cells) > 0,
            formula_cells=all_formula_cells,
            warnings=warnings,
        )

    finally:
        wb_formula.close()
        wb_values.close()


def get_worksheet_preview(
    file_path: Path | str,
    sheet_name: str,
    max_rows: int = 20,
    max_cols: int = 15,
    max_size_mb: int = DEFAULT_MAX_FILE_SIZE_MB,
) -> list[list[Any]]:
    """
    Extract a preview grid (2D matrix) of cell values from a worksheet.

    Reads cached values (data_only=True) so users see actual data rather
    than formula strings.

    Args:
        file_path: Path to the workbook.
        sheet_name: Name of the worksheet to preview.
        max_rows: Maximum rows to return.
        max_cols: Maximum columns to return.
        max_size_mb: Maximum allowed file size.

    Returns:
        A list of rows, where each row is a list of cell values.

    Raises:
        WorksheetNotFoundError: If the requested sheet does not exist.
        CorruptWorkbookError: If the workbook cannot be parsed.
    """
    validate_xlsx_file(file_path, max_size_mb=max_size_mb)
    path = Path(file_path).resolve()

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CorruptWorkbookError(
            f"Failed to open workbook for preview: {exc}"
        ) from exc

    try:
        if sheet_name not in wb.sheetnames:
            raise WorksheetNotFoundError(
                f"Worksheet '{sheet_name}' not found in workbook. Available sheets: {wb.sheetnames}"
            )

        ws = wb[sheet_name]
        grid: list[list[Any]] = []

        for row_idx, row in enumerate(
            ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True),
            start=1,
        ):
            grid.append(list(row))

        return grid
    finally:
        wb.close()
