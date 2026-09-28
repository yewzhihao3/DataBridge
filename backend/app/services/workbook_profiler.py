"""Bounded read-only analysis; values stay internal, metadata is returned to clients."""
from dataclasses import dataclass
from datetime import date, datetime
import re
import openpyxl
from openpyxl.utils import get_column_letter
from app.schemas.suggestion import SheetProfile
from app.services.workbook_inspector import validate_xlsx_file, CorruptWorkbookError

MAX_ROWS = 2000
MAX_COLUMNS = 100
MAX_SHEETS = 30
CURRENCIES = {"USD", "MYR", "SGD", "EUR", "GBP", "AUD", "CAD", "JPY", "CNY", "INR", "NZD", "HKD"}


def normalize(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower().replace("#", " no ")).strip()


def kind(value, formula=False):
    if formula:
        return "formula"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    text = str(value).strip()
    if re.fullmatch(r"(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/][A-Za-z]{3}[-/]\d{4}|\d{1,2}[-/]\d{1,2}[-/]\d{4})", text):
        return "date"
    if text in CURRENCIES or re.match(r"^(?:[$€£¥]|(?:USD|MYR|SGD|EUR|GBP)\s)\s*\d", text):
        return "currency"
    if re.fullmatch(r"[-+]?\d[\d,]*(?:\.\d+)?%?", text):
        return "number"
    return "text"


@dataclass
class Cell:
    row: int
    col: int
    ref: str
    value: object
    kind: str
    end_row: int
    end_col: int


@dataclass
class Profile:
    metadata: SheetProfile
    cells: dict[tuple[int, int], Cell]


def profile_workbook(path, max_size_mb=10):
    validate_xlsx_file(path, max_size_mb)
    try:
        wb = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_vba=False)
    except Exception as exc:
        raise CorruptWorkbookError("Could not analyze this workbook.") from exc
    profiles = []
    try:
        # Report every sheet even if the scan budget is exhausted.
        for index, ws in enumerate(wb.worksheets):
            rows = min(ws.max_row, MAX_ROWS) if index < MAX_SHEETS else 0
            cols = min(ws.max_column, MAX_COLUMNS) if rows else 0
            merges = {(r.min_row, r.min_col): r for r in ws.merged_cells.ranges}
            cells = {}
            row_cells = {}
            for row in ws.iter_rows(max_row=rows, max_col=cols) if rows else []:
                for c in row:
                    if c.value is None or str(c.value).strip() == "":
                        continue
                    merge = merges.get((c.row, c.column))
                    cell = Cell(c.row, c.column, c.coordinate, c.value, kind(c.value, c.data_type == "f"),
                                merge.max_row if merge else c.row, merge.max_col if merge else c.column)
                    cells[c.row, c.column] = cell
                    row_cells.setdefault(c.row, []).append(cell)
            numeric = [get_column_letter(c) for c in range(1, cols + 1)
                       if sum(x.col == c and x.kind in {"number", "formula"} for x in cells.values()) >= 2]
            headers = [r for r, values in row_cells.items() if len(values) >= 2
                       and all(x.kind == "text" for x in values)
                       and len(row_cells.get(r + 1, [])) >= 2][:30]
            profiles.append(Profile(SheetProfile(
                name=ws.title, is_hidden=ws.sheet_state != "visible", used_range=ws.calculate_dimension(),
                scanned_rows=rows, scanned_columns=cols,
                truncated=ws.max_row > rows or ws.max_column > cols,
                non_empty_count=len(cells), likely_header_rows=headers,
                blank_rows=[r for r in range(1, rows + 1) if r not in row_cells][:100],
                numeric_columns=numeric,
                date_cells=[c.ref for c in cells.values() if c.kind == "date"][:100],
                currency_cells=[c.ref for c in cells.values() if c.kind == "currency"][:100],
                formula_cells=[c.ref for c in cells.values() if c.kind == "formula"][:100],
                merged_ranges=[str(r) for r in ws.merged_cells.ranges][:100],
            ), cells))
        return profiles
    finally:
        wb.close()
