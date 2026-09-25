import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from decimal import Decimal
import openpyxl
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker, selectinload

from app.models import ImportBatch, InvoiceLineItem, InvoiceRecord, SourceFile, Template, TemplateFieldMapping
from app.services.extractor import extract_from_workbook
from app.services.validator import validate_extraction
import json

def run_qa():
    print("=== STARTING REALISTIC INVOICE MANUAL QA ===")
    
    # 1. Create a realistic Excel workbook in memory
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TaxInvoice"
    
    # Header cells
    ws["B2"] = "INV-2026-ACME-883"
    ws["B3"] = "2026-09-26"
    ws["B4"] = "Acme Heavy Industries LLC"
    ws["E2"] = "Global Transport Solutions"
    ws["E3"] = "USD"
    ws["E4"] = "NET 30"
    ws["E18"] = 6250.00   # Subtotal
    ws["E19"] = 500.00    # Tax
    ws["E20"] = 6750.00   # Total
    
    # Table headers at row 8
    ws["A8"] = "Item Description"
    ws["B8"] = "SKU"
    ws["C8"] = "Qty"
    ws["D8"] = "Unit Price"
    ws["E8"] = "Tax Rate"
    ws["F8"] = "Tax Amount"
    ws["G8"] = "Line Total"
    
    # 4 Realistic Line Items (rows 9..12)
    # Including products with potential footer words ("Total Care...", "Tax Preparation...")
    items = [
        ("Precision Ball Bearings 50mm", "SKU-BEAR-01", 100, 25.00, "8%", 200.00, 2500.00),
        ("Total Care Cleaning Kit", "SKU-CARE-09", 10, 75.00, "8%", 60.00, 750.00),
        ("Tax Preparation Service", "SKU-SRV-99", 2, 500.00, "8%", 80.00, 1000.00),
        ("High Pressure Hydraulic Valve", "SKU-VALV-44", 20, 100.00, "8%", 160.00, 2000.00),
    ]
    
    for idx, (desc, sku, qty, price, rate, tax, tot) in enumerate(items, start=9):
        ws[f"A{idx}"] = desc
        ws[f"B{idx}"] = sku
        ws[f"C{idx}"] = qty
        ws[f"D{idx}"] = price
        ws[f"E{idx}"] = rate
        ws[f"F{idx}"] = tax
        ws[f"G{idx}"] = tot
        
    # Footer rows at rows 13..15
    ws["A13"] = "Subtotal"
    ws["G13"] = 6250.00
    
    ws["A14"] = "Tax"
    ws["G14"] = 500.00
    
    ws["A15"] = "TOTAL"
    ws["G15"] = 6750.00
    
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()
    
    # 2. Build mock Template with Header and Line Item mappings
    template = Template(
        id=999,
        name="Acme Composite Invoice Template",
        template_type="invoice",
        file_type="xlsx",
        worksheet="TaxInvoice",
        header_row=8,
        data_start_row=9,
        field_mappings=[
            # Header
            TemplateFieldMapping(field_name="invoice_number", mapping_group="header", mapping_type="cell", cell_ref="B2", is_required=True, data_type="text"),
            TemplateFieldMapping(field_name="invoice_date", mapping_group="header", mapping_type="cell", cell_ref="B3", is_required=True, data_type="date"),
            TemplateFieldMapping(field_name="supplier_name", mapping_group="header", mapping_type="cell", cell_ref="B4", is_required=True, data_type="text"),
            TemplateFieldMapping(field_name="customer_name", mapping_group="header", mapping_type="cell", cell_ref="E2", is_required=False, data_type="text"),
            TemplateFieldMapping(field_name="currency", mapping_group="header", mapping_type="cell", cell_ref="E3", is_required=False, data_type="text"),
            TemplateFieldMapping(field_name="subtotal_amount", mapping_group="header", mapping_type="cell", cell_ref="E18", is_required=False, data_type="decimal"),
            TemplateFieldMapping(field_name="tax_amount", mapping_group="header", mapping_type="cell", cell_ref="E19", is_required=False, data_type="decimal"),
            TemplateFieldMapping(field_name="total_amount", mapping_group="header", mapping_type="cell", cell_ref="E20", is_required=True, data_type="decimal"),
            # Line Items (including custom field 'sku_code')
            TemplateFieldMapping(field_name="description", mapping_group="line_item", mapping_type="column", column_ref="A", is_required=True, data_type="text"),
            TemplateFieldMapping(field_name="sku_code", mapping_group="line_item", mapping_type="column", column_ref="B", is_required=False, data_type="text"),
            TemplateFieldMapping(field_name="quantity", mapping_group="line_item", mapping_type="column", column_ref="C", is_required=True, data_type="decimal"),
            TemplateFieldMapping(field_name="unit_price", mapping_group="line_item", mapping_type="column", column_ref="D", is_required=True, data_type="decimal"),
            TemplateFieldMapping(field_name="tax_rate", mapping_group="line_item", mapping_type="column", column_ref="E", is_required=False, data_type="decimal"),
            TemplateFieldMapping(field_name="tax_amount", mapping_group="line_item", mapping_type="column", column_ref="F", is_required=False, data_type="decimal"),
            TemplateFieldMapping(field_name="amount", mapping_group="line_item", mapping_type="column", column_ref="G", is_required=True, data_type="decimal"),
        ]
    )
    
    # 3. Test Extraction
    wb_formula = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    wb_values = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    extract_res = extract_from_workbook(wb_formula=wb_formula, wb_values=wb_values, template=template)
    
    print(f"Extraction result: has_line_items={extract_res.has_line_items}, line_items_count={len(extract_res.line_items)}")
    assert extract_res.has_line_items is True
    assert len(extract_res.line_items) == 4, f"Expected 4 line items, got {len(extract_res.line_items)}"
    
    extracted_descs = [
        next(f.normalized_value for f in li.fields if f.field_name == "description")
        for li in extract_res.line_items
    ]
    print(f"Extracted line item descriptions: {extracted_descs}")
    assert extracted_descs == [
        "Precision Ball Bearings 50mm",
        "Total Care Cleaning Kit",
        "Tax Preparation Service",
        "High Pressure Hydraulic Valve",
    ]
    
    # 4. Test Validation
    val_report = validate_extraction(extract_res)
    print(f"Validation report: is_valid={val_report.is_valid_for_import}, errors={val_report.error_count}, warnings={val_report.warning_count}")
    assert val_report.is_valid_for_import is True
    assert val_report.error_count == 0
    
    # 5. Test Persistence in SQLite in-memory DB
    engine = create_engine("sqlite:///:memory:")
    from app.database import Base
    Base.metadata.create_all(engine)
    
    SessionLocal = sessionmaker(bind=engine)
    db = SessionLocal()
    
    sf = SourceFile(
        id=1,
        stored_filename="test.xlsx",
        original_name="test.xlsx",
        file_size=len(file_bytes),
        checksum="abc123"
    )
    db.add(sf)
    db.add(template)
    db.flush()
    
    # Create batch
    batch = ImportBatch(
        source_file_id=sf.id,
        template_id=template.id,
        status="imported",
        record_count=1,
        warning_count=0
    )
    db.add(batch)
    db.flush()
    
    norm_data = val_report.normalized_data
    inv_rec = InvoiceRecord(
        batch_id=batch.id,
        company_name=norm_data.get("supplier_name", "Acme Heavy Industries LLC"),
        invoice_number=norm_data.get("invoice_number", "INV-2026-ACME-883"),
        invoice_date=norm_data.get("invoice_date"),
        total_amount=norm_data.get("total_amount"),
        currency=norm_data.get("currency", "USD"),
        source_worksheet=template.worksheet,
        raw_data=json.dumps({"test": 1}),
        custom_fields={"subtotal_amount": "6250.00", "tax_amount": "500.00"}
    )
    db.add(inv_rec)
    db.flush()
    
    # Add child line items
    for lr in val_report.line_item_reports:
        li_norm = lr.normalized_data
        li = InvoiceLineItem(
            invoice_id=inv_rec.id,
            source_row_number=lr.source_row_number,
            description=li_norm.get("description"),
            quantity=li_norm.get("quantity"),
            unit_price=li_norm.get("unit_price"),
            tax_rate=li_norm.get("tax_rate"),
            tax_amount=li_norm.get("tax_amount"),
            amount=li_norm.get("amount"),
            custom_fields={"sku_code": li_norm.get("sku_code")} if "sku_code" in li_norm else None,
            raw_data=json.dumps({k: str(v) for k, v in li_norm.items()})
        )
        db.add(li)
        
    db.commit()
    
    # 6. Verify persisted entities
    persisted_inv = db.query(InvoiceRecord).filter(InvoiceRecord.id == inv_rec.id).options(selectinload(InvoiceRecord.line_items)).one()
    print("\n=== PERSISTENCE VERIFICATION ===")
    print(f"Invoice Record: ID={persisted_inv.id}, Number={persisted_inv.invoice_number}, Total={persisted_inv.total_amount}, Currency={persisted_inv.currency}")
    print(f"Child Line Items Count: {len(persisted_inv.line_items)}")
    
    assert len(persisted_inv.line_items) == 4
    for idx, li in enumerate(persisted_inv.line_items, 1):
        print(f"  Line {idx}: Row {li.source_row_number} | {li.description} | Qty: {li.quantity} | Unit: {li.unit_price} | Tax Rate: {li.tax_rate}% | Tax: {li.tax_amount} | Total: {li.amount} | Custom: {li.custom_fields}")
        assert li.source_row_number in [9, 10, 11, 12]
        assert li.custom_fields is not None and "sku_code" in li.custom_fields
        
    print("\nALL QA CHECKS PASSED PERFECTLY!")

if __name__ == "__main__":
    run_qa()
