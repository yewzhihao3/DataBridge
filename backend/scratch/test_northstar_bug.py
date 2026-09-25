import sys
from pathlib import Path
from decimal import Decimal
import openpyxl
import json
from datetime import date

backend_dir = Path(r"c:\Users\Bruce Yew\Desktop\Project\Databridge\backend")
sys.path.insert(0, str(backend_dir))

from app.models import ImportBatch, InvoiceRecord, InvoiceLineItem, Template, TemplateFieldMapping, SourceFile
from app.services.extractor import extract_from_workbook
from app.services.validator import validate_extraction, ValidationConfig
from app.constants import CANONICAL_INVOICE_FIELDS, CANONICAL_LINE_ITEM_FIELDS
from app.services.normalizer import normalize_date, normalize_decimal
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app.database import Base

CANONICAL_INVOICE_ALIASES: dict[str, tuple[str, ...]] = {
    "company_name": ("company_name", "company", "supplier", "supplier_name", "vendor", "vendor_name", "organization", "bill_from"),
    "invoice_number": ("invoice_number", "invoice_num", "invoice_no", "inv_no", "inv_num", "inv_number", "invoice_id"),
    "invoice_date": ("invoice_date", "date", "inv_date", "issue_date", "billing_date"),
    "total_amount": ("total_amount", "amount", "total", "total_amt", "grand_total", "invoice_total", "total_due", "amount_due", "net_amount"),
    "currency": ("currency", "curr", "currency_code"),
}

CANONICAL_LINE_ITEM_ALIASES: dict[str, tuple[str, ...]] = {
    "description": ("description", "desc", "item", "product", "item_description", "product_description", "details", "particulars"),
    "quantity": ("quantity", "qty", "count", "units"),
    "unit_price": ("unit_price", "price", "rate", "unit_rate", "price_unit"),
    "tax_rate": ("tax_rate", "tax_pct", "tax_percent", "tax_rate_pct", "vat_rate", "gst_rate"),
    "tax_amount": ("tax_amount", "tax", "tax_amt", "vat_amount", "gst_amount"),
    "amount": ("amount", "line_total", "total", "line_amount", "subtotal", "extended_amount"),
}

_norm_canonical_set = {
    str(alias).strip().lower().replace(" ", "_").replace("-", "_")
    for key, aliases in CANONICAL_INVOICE_ALIASES.items()
    for alias in (key, *aliases)
}

_norm_line_canonical_set = {
    str(alias).strip().lower().replace(" ", "_").replace("-", "_")
    for key, aliases in CANONICAL_LINE_ITEM_ALIASES.items()
    for alias in (key, *aliases)
}

from typing import Any

def _resolve_row_canonical(norm_data: dict, canonical_key: str, aliases: tuple[str, ...]) -> Any:
    for k in (canonical_key, *aliases):
        if k in norm_data and norm_data[k] is not None:
            return norm_data[k]
    _norm_lookup = {
        str(_k).strip().lower().replace(" ", "_").replace("-", "_"): _v
        for _k, _v in norm_data.items()
        if _v is not None
    }
    for k in (canonical_key, *aliases):
        norm_k = str(k).strip().lower().replace(" ", "_").replace("-", "_")
        if norm_k in _norm_lookup:
            return _norm_lookup[norm_k]
    return None

def run_confirm_simulation(template):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"

    ws["A1"] = "NORTHSTAR OFFICE SOLUTIONS SDN. BHD."
    ws["F5"] = "NOS-2026-0947"
    ws["F6"] = "20-Sep-2026"
    ws["F7"] = "20-Oct-2026"
    ws["F8"] = "MYR"

    ws["A10"] = "Description"
    ws["B10"] = "Qty"
    ws["C10"] = "Unit Price"
    ws["D10"] = "Tax Rate"
    ws["E10"] = "Amount"

    ws["A11"] = "Ergonomic Office Chair — Model E7"
    ws["B11"] = 6
    ws["C11"] = 689.00
    ws["D11"] = "8%"
    ws["E11"] = 4134.00

    ws["A12"] = "Adjustable Monitor Arm — Dual Display"
    ws["B12"] = 6
    ws["C12"] = 249.00
    ws["D12"] = "8%"
    ws["E12"] = 1494.00

    ws["A13"] = "Wireless Keyboard & Mouse Set"
    ws["B13"] = 6
    ws["C13"] = 129.90
    ws["D13"] = "8%"
    ws["E13"] = 779.40

    ws["A14"] = "Cable Management Tray"
    ws["B14"] = 6
    ws["C14"] = 45.00
    ws["D14"] = "8%"
    ws["E14"] = 270.00

    ws["D16"] = "Subtotal"
    ws["E16"] = 6677.40

    ws["D17"] = "Tax (8%)"
    ws["E17"] = 534.192

    ws["D18"] = "TOTAL"
    ws["E18"] = 7211.592

    import io
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    wb_formula = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    wb_values = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    ext_res = extract_from_workbook(wb_formula, wb_values, template)
    field_to_target = {m.field_name: m.target_field for m in template.field_mappings if m.target_field}
    val_report = validate_extraction(ext_res, field_to_target=field_to_target)

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    db = Session()

    sf = SourceFile(id=1, stored_filename="northstar.xlsx", original_name="northstar.xlsx", file_size=len(file_bytes), checksum="123")
    db.add(sf)
    db.add(template)
    db.flush()

    def _coerce_date(val):
        if val is None:
            return None
        if isinstance(val, date):
            return val
        if isinstance(val, str):
            try:
                return normalize_date(val)
            except Exception:
                return None
        return None

    def _coerce_decimal(val):
        if val is None:
            return None
        if isinstance(val, Decimal):
            return val
        try:
            return normalize_decimal(val)
        except Exception:
            return None

    norm_data = val_report.normalized_data
    comp_name = _resolve_row_canonical(norm_data, "company_name", CANONICAL_INVOICE_ALIASES["company_name"]) or "Unknown Company"
    inv_num = _resolve_row_canonical(norm_data, "invoice_number", CANONICAL_INVOICE_ALIASES["invoice_number"]) or "Unknown Invoice"
    inv_date = _coerce_date(_resolve_row_canonical(norm_data, "invoice_date", CANONICAL_INVOICE_ALIASES["invoice_date"]))
    tot_amt = _coerce_decimal(_resolve_row_canonical(norm_data, "total_amount", CANONICAL_INVOICE_ALIASES["total_amount"]))
    curr = _resolve_row_canonical(norm_data, "currency", CANONICAL_INVOICE_ALIASES["currency"])

    custom_fields = {}
    for key, val in norm_data.items():
        norm_k = str(key).strip().lower().replace(" ", "_").replace("-", "_")
        if key not in CANONICAL_INVOICE_FIELDS and norm_k not in _norm_canonical_set:
            custom_fields[key] = str(val)

    batch = ImportBatch(source_file_id=sf.id, template_id=template.id, status="imported", record_count=1, warning_count=0)
    inv_rec = InvoiceRecord(
        company_name=str(comp_name),
        invoice_number=str(inv_num),
        invoice_date=inv_date,
        total_amount=tot_amt,
        currency=str(curr) if curr else None,
        source_worksheet="Invoice",
        raw_data="{}",
        custom_fields=custom_fields if custom_fields else None
    )

    if val_report.has_line_items and val_report.line_item_reports:
        for lr in val_report.line_item_reports:
            li_norm = lr.normalized_data
            li_desc = _resolve_row_canonical(li_norm, "description", CANONICAL_LINE_ITEM_ALIASES["description"])
            li_qty = _coerce_decimal(_resolve_row_canonical(li_norm, "quantity", CANONICAL_LINE_ITEM_ALIASES["quantity"]))
            li_price = _coerce_decimal(_resolve_row_canonical(li_norm, "unit_price", CANONICAL_LINE_ITEM_ALIASES["unit_price"]))
            li_tax_rate = _coerce_decimal(_resolve_row_canonical(li_norm, "tax_rate", CANONICAL_LINE_ITEM_ALIASES["tax_rate"]))
            li_tax_amt = _coerce_decimal(_resolve_row_canonical(li_norm, "tax_amount", CANONICAL_LINE_ITEM_ALIASES["tax_amount"]))
            li_amt = _coerce_decimal(_resolve_row_canonical(li_norm, "amount", CANONICAL_LINE_ITEM_ALIASES["amount"]))

            li_custom = {}
            for key, val in li_norm.items():
                norm_k = str(key).strip().lower().replace(" ", "_").replace("-", "_")
                if key not in CANONICAL_LINE_ITEM_FIELDS and norm_k not in _norm_line_canonical_set:
                    li_custom[key] = str(val)

            line_item = InvoiceLineItem(
                source_row_number=lr.source_row_number,
                description=str(li_desc) if li_desc is not None else None,
                quantity=li_qty,
                unit_price=li_price,
                tax_rate=li_tax_rate,
                tax_amount=li_tax_amt,
                amount=li_amt,
                custom_fields=li_custom if li_custom else None,
                raw_data="{}"
            )
            inv_rec.line_items.append(line_item)

    batch.invoice_records.append(inv_rec)
    db.add(batch)
    db.commit()

    db.refresh(inv_rec)
    print("=== PERSISTED INVOICE RECORD ===")
    print(f"company_name: {inv_rec.company_name}")
    print(f"invoice_number: {inv_rec.invoice_number}")
    print(f"invoice_date: {inv_rec.invoice_date}")
    print(f"total_amount: {inv_rec.total_amount}")
    print(f"currency: {inv_rec.currency}")
    print(f"custom_fields: {inv_rec.custom_fields}")
    print("=== PERSISTED LINE ITEMS ===")
    for li in inv_rec.line_items:
        print(f"Row {li.source_row_number}: desc='{li.description}', qty={li.quantity}, price={li.unit_price}, tax_rate={li.tax_rate}, tax_amt={li.tax_amount}, amt={li.amount}, custom={li.custom_fields}")

# Case A: Template with target_fields set
t1 = Template(
    id=10, name="T1", template_type="invoice", worksheet="Invoice", header_row=10, data_start_row=11,
    field_mappings=[
        TemplateFieldMapping(field_name="Supplier", target_field="company_name", mapping_group="header", mapping_type="cell", cell_ref="A1"),
        TemplateFieldMapping(field_name="Invoice No", target_field="invoice_number", mapping_group="header", mapping_type="cell", cell_ref="F5"),
        TemplateFieldMapping(field_name="Invoice Date", target_field="invoice_date", mapping_group="header", mapping_type="cell", cell_ref="F6"),
        TemplateFieldMapping(field_name="Due Date", target_field="due_date", mapping_group="header", mapping_type="cell", cell_ref="F7"),
        TemplateFieldMapping(field_name="Currency", target_field="currency", mapping_group="header", mapping_type="cell", cell_ref="F8"),
        TemplateFieldMapping(field_name="Subtotal", target_field="sub_total", mapping_group="header", mapping_type="cell", cell_ref="E16"),
        TemplateFieldMapping(field_name="Tax Amount", target_field="tax_amount", mapping_group="header", mapping_type="cell", cell_ref="E17"),
        TemplateFieldMapping(field_name="Total Amount", target_field="total_amount", mapping_group="header", mapping_type="cell", cell_ref="E18"),

        TemplateFieldMapping(field_name="Description", target_field="description", mapping_group="line_item", mapping_type="column", column_ref="A"),
        TemplateFieldMapping(field_name="Qty", target_field="quantity", mapping_group="line_item", mapping_type="column", column_ref="B"),
        TemplateFieldMapping(field_name="Unit Price", target_field="unit_price", mapping_group="line_item", mapping_type="column", column_ref="C"),
        TemplateFieldMapping(field_name="Tax Rate", target_field="tax_rate", mapping_group="line_item", mapping_type="column", column_ref="D"),
        TemplateFieldMapping(field_name="Amount", target_field="amount", mapping_group="line_item", mapping_type="column", column_ref="E"),
    ]
)

# Case B: Template where target_field is NOT set, relying on canonical aliases
t2 = Template(
    id=20, name="T2", template_type="invoice", worksheet="Invoice", header_row=10, data_start_row=11,
    field_mappings=[
        TemplateFieldMapping(field_name="company_name", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="A1"),
        TemplateFieldMapping(field_name="invoice_number", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="F5"),
        TemplateFieldMapping(field_name="invoice_date", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="F6"),
        TemplateFieldMapping(field_name="due_date", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="F7"),
        TemplateFieldMapping(field_name="currency", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="F8"),
        TemplateFieldMapping(field_name="sub_total", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="E16"),
        TemplateFieldMapping(field_name="tax_amount", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="E17"),
        TemplateFieldMapping(field_name="Total Invoice Amount", target_field=None, mapping_group="header", mapping_type="cell", cell_ref="E18"),

        TemplateFieldMapping(field_name="Item Description", target_field=None, mapping_group="line_item", mapping_type="column", column_ref="A"),
        TemplateFieldMapping(field_name="Quantity", target_field=None, mapping_group="line_item", mapping_type="column", column_ref="B"),
        TemplateFieldMapping(field_name="Unit Price", target_field=None, mapping_group="line_item", mapping_type="column", column_ref="C"),
        TemplateFieldMapping(field_name="Tax Rate", target_field=None, mapping_group="line_item", mapping_type="column", column_ref="D"),
        TemplateFieldMapping(field_name="Line Total", target_field=None, mapping_group="line_item", mapping_type="column", column_ref="E"),
    ]
)

print("--- RUNNING CASE A ---")
run_confirm_simulation(t1)

print("\n--- RUNNING CASE B ---")
run_confirm_simulation(t2)
