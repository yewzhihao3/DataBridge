"""
tests/integration/test_api_multi_record_imports.py — Integration tests for multi-record imports.
"""

from pathlib import Path
import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy.orm import Session

from app.models.invoice import ImportBatch, InvoiceRecord
from app.models.template import Template, TemplateFieldMapping


@pytest.fixture
def multi_row_excel_file(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # Header Row 1
    ws["A1"] = "Supplier"
    ws["B1"] = "Invoice No"
    ws["C1"] = "Date"
    ws["D1"] = "Amount"
    ws["E1"] = "PO Number"

    # Row 2
    ws["A2"] = "ABC Sdn Bhd"
    ws["B2"] = "INV-001"
    ws["C2"] = "2026-09-20"
    ws["D2"] = 1200.00
    ws["E2"] = "PO-101"

    # Row 3
    ws["A3"] = "XYZ Sdn Bhd"
    ws["B3"] = "INV-002"
    ws["C3"] = "2026-09-21"
    ws["D3"] = 850.50
    ws["E3"] = "PO-102"

    file_path = tmp_path / "multi_invoices.xlsx"
    wb.save(file_path)
    return file_path


@pytest.fixture
def multi_row_template(db_session: Session) -> Template:
    template = Template(
        name="Multi Row Invoice Template",
        file_type="xlsx",
        worksheet="Invoices",
        header_row=1,
        data_start_row=2,
    )
    template.field_mappings = [
        TemplateFieldMapping(field_name="supplier", target_field="company_name", mapping_type="column", column_ref="A", is_required=True, data_type="text"),
        TemplateFieldMapping(field_name="inv_no", target_field="invoice_number", mapping_type="column", column_ref="B", is_required=True, data_type="text"),
        TemplateFieldMapping(field_name="inv_date", target_field="invoice_date", mapping_type="column", column_ref="C", is_required=False, data_type="date"),
        TemplateFieldMapping(field_name="amount", target_field="total_amount", mapping_type="column", column_ref="D", is_required=False, data_type="decimal"),
        TemplateFieldMapping(field_name="po_number", target_field="purchase_order_number", mapping_type="column", column_ref="E", is_required=False, data_type="text"),
    ]
    db_session.add(template)
    db_session.commit()
    db_session.refresh(template)
    return template


@pytest.mark.asyncio
async def test_multi_record_extract_preview_api(
    client: AsyncClient,
    db_session: Session,
    multi_row_excel_file: Path,
    multi_row_template: Template,
) -> None:
    # 1. Upload source file
    with open(multi_row_excel_file, "rb") as f:
        res = await client.post("/api/v1/files/upload", files={"file": ("multi_invoices.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res.status_code == 201
    file_id = res.json()["file_id"]

    # 2. Call /extract preview
    res_preview = await client.post("/api/v1/imports/extract", json={"file_id": file_id, "template_id": multi_row_template.id})
    assert res_preview.status_code == 200
    data = res_preview.json()

    assert data["is_multi_record"] is True
    assert data["record_count"] == 2
    assert len(data["records"]) == 2
    assert data["records"][0]["source_row_number"] == 2
    assert data["records"][0]["normalized_data"]["company_name"] == "ABC Sdn Bhd"
    assert data["records"][0]["normalized_data"]["purchase_order_number"] == "PO-101"
    assert data["records"][1]["source_row_number"] == 3
    assert data["records"][1]["normalized_data"]["company_name"] == "XYZ Sdn Bhd"
    assert data["records"][1]["normalized_data"]["purchase_order_number"] == "PO-102"


@pytest.mark.asyncio
async def test_multi_record_confirm_persistence_and_custom_field_isolation(
    client: AsyncClient,
    db_session: Session,
    multi_row_excel_file: Path,
    multi_row_template: Template,
) -> None:
    # 1. Upload file
    with open(multi_row_excel_file, "rb") as f:
        res = await client.post("/api/v1/files/upload", files={"file": ("multi_invoices.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    file_id = res.json()["file_id"]

    # 2. Confirm import
    res_confirm = await client.post("/api/v1/imports/confirm", json={"file_id": file_id, "template_id": multi_row_template.id})
    assert res_confirm.status_code == 201
    batch_data = res_confirm.json()

    assert batch_data["record_count"] == 2
    assert len(batch_data["invoice_records"]) == 2

    rec1 = batch_data["invoice_records"][0]
    rec2 = batch_data["invoice_records"][1]

    assert rec1["company_name"] == "ABC Sdn Bhd"
    assert rec1["invoice_number"] == "INV-001"
    assert rec1["source_row_number"] == 2
    assert rec1["custom_fields"] == {"purchase_order_number": "PO-101"}

    assert rec2["company_name"] == "XYZ Sdn Bhd"
    assert rec2["invoice_number"] == "INV-002"
    assert rec2["source_row_number"] == 3
    assert rec2["custom_fields"] == {"purchase_order_number": "PO-102"}

    # Verify DB state
    db_batch = db_session.query(ImportBatch).filter(ImportBatch.id == batch_data["id"]).first()
    assert db_batch is not None
    assert db_batch.record_count == 2
    assert len(db_batch.invoice_records) == 2
