"""M11 semantic regressions; suggestions must remain review-only."""
import io
from types import SimpleNamespace
import openpyxl
import pytest
from app.services.workbook_profiler import profile_workbook
from app.services.structure_detector import analyze_sheet
from app.services.template_matcher import match_templates
from tests.integration.test_m75_northstar_persistence_audit import build_northstar_workbook_bytes


def northstar(labels=True):
    wb = openpyxl.load_workbook(io.BytesIO(build_northstar_workbook_bytes()))
    ws = wb.active
    ws.merge_cells('A1:F1')
    if labels:
        for row, text in [(5, 'Invoice No'), (6, 'Invoice Date'), (7, 'Due Date'), (8, 'Currency')]:
            ws[f'E{row}'] = text
    return wb


def analyze(wb, tmp_path):
    path = tmp_path / 'arbitrary-name.xlsx'
    wb.save(path)
    return analyze_sheet(profile_workbook(path)[0])


async def upload(client, wb):
    stream = io.BytesIO()
    wb.save(stream)
    response = await client.post('/api/v1/files/upload', files={'file': ('sample.xlsx', stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 201
    return response.json()['file_id']


@pytest.mark.parametrize('labels', [True, False])
def test_northstar_detection(tmp_path, labels):
    result = analyze(northstar(labels), tmp_path)
    assert result.suggested_template_type == 'invoice'
    assert result.confidence in {'high', 'medium'}
    assert {m.target_field: m.cell_ref for m in result.mappings} == {
        'company_name': 'A1', 'invoice_number': 'F5', 'invoice_date': 'F6', 'due_date': 'F7',
        'currency': 'F8', 'sub_total': 'E16', 'tax_amount': 'E17', 'total_amount': 'E18',
    }
    assert (result.table.header_row, result.table.data_start_row, result.table.data_end_row) == (10, 11, 14)
    assert {m.target_field: m.column_ref for m in result.table.columns} == {
        'description': 'A', 'quantity': 'B', 'unit_price': 'C', 'tax_rate': 'D', 'amount': 'E',
    }


@pytest.mark.parametrize('alias', ['Invoice No', 'Invoice #', 'Inv No', 'Invoice Number', 'Document No', 'INVOICE_NO:'])
def test_invoice_aliases(tmp_path, alias):
    wb = northstar()
    wb.active['E5'] = alias
    assert next(m for m in analyze(wb, tmp_path).mappings if m.target_field == 'invoice_number').cell_ref == 'F5'


@pytest.mark.parametrize('alias', ['Grand Total', 'Invoice Total', 'Amount Due', 'TOTAL'])
def test_total_aliases(tmp_path, alias):
    wb = northstar()
    wb.active['D18'] = alias
    assert next(m for m in analyze(wb, tmp_path).mappings if m.target_field == 'total_amount').cell_ref == 'E18'


def test_merged_label_and_value_below(tmp_path):
    wb = northstar()
    ws = wb.active
    ws['E5'] = None
    ws['F5'] = None
    ws['A4'] = 'Document No'
    ws.merge_cells('A4:C4')
    ws['D4'] = 'DOC-1'
    ws['E6'] = None
    ws['F6'] = None
    ws['B6'] = 'Invoice Date'
    ws['B7'] = '2026-09-20'
    fields = {m.target_field: m.cell_ref for m in analyze(wb, tmp_path).mappings}
    assert fields['invoice_number'] == 'D4'
    assert fields['invoice_date'] == 'B7'


@pytest.mark.parametrize('description', ['Total Care Cleaning Kit', 'Tax Preparation Service', 'GST Compliance Software'])
def test_product_words_do_not_end_table(tmp_path, description):
    wb = northstar()
    wb.active['A12'] = description
    assert analyze(wb, tmp_path).table.data_end_row == 14


def dataset():
    wb = openpyxl.Workbook()
    wb.active.title = 'Records'
    wb.active.append(['Company Name', 'Invoice Number', 'Invoice Date', 'Currency', 'Total Amount'])
    wb.active.append(['Acme', 'A-1', '2026-09-01', 'MYR', 100])
    wb.active.append(['Beta', 'A-2', '2026-09-02', 'MYR', 200])
    return wb


def test_dataset_headers(tmp_path):
    result = analyze(dataset(), tmp_path)
    assert result.suggested_template_type == 'dataset'
    assert result.confidence == 'high'
    assert result.table.header_row == 1
    assert result.table.data_start_row == 2
    assert {m.target_field: m.column_ref for m in result.mappings} == {
        'company_name': 'A', 'invoice_number': 'B', 'invoice_date': 'C', 'currency': 'D', 'total_amount': 'E'}


def test_ambiguous_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(['Notes', 'Invoice'])
    result = analyze(wb, tmp_path)
    assert result.confidence == 'low'
    assert result.suggested_template_type is None
    assert result.mappings == []


def test_generic_dataset_requires_consistent_rows(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(['Location', 'Reading', 'Observed'])
    for value in range(3):
        wb.active.append(['Warehouse', value, '2026-09-01'])
    result = analyze(wb, tmp_path)
    assert result.suggested_template_type == 'dataset'
    assert result.confidence == 'medium'
    assert result.mappings[1].target_field == 'custom_reading'
    assert result.mappings[1].data_type == 'decimal'
    wb.active['B3'] = 'Notes'
    assert analyze(wb, tmp_path).confidence == 'low'


def test_adjacent_labels_with_values_below(tmp_path):
    wb = northstar()
    ws = wb.active
    ws['E5'] = ws['F5'] = ws['E6'] = ws['F6'] = None
    ws['A4'], ws['B4'] = 'Invoice No', 'Invoice Date'
    ws['A5'], ws['B5'] = 'DOC-123', '2026-09-20'
    fields = {m.target_field: m.cell_ref for m in analyze(wb, tmp_path).mappings}
    assert fields['invoice_number'] == 'A5'
    assert fields['invoice_date'] == 'B5'


def test_currency_amount_is_not_currency_code(tmp_path):
    wb = northstar()
    wb.active['F8'] = 'USD 100'
    assert not any(m.target_field == 'currency' for m in analyze(wb, tmp_path).mappings)


def test_profile_metadata_bounded_and_non_mutating(tmp_path):
    wb = northstar()
    wb.active['E14'] = '=B14*C14'
    wb.active['A2002'] = 'outside scan'
    path = tmp_path / 'profile.xlsx'
    wb.save(path)
    before = path.read_bytes()
    profile = profile_workbook(path)[0].metadata
    assert path.read_bytes() == before
    assert profile.truncated
    assert profile.scanned_rows == 2000
    assert profile.formula_cells == ['E14']
    assert 'A1:F1' in profile.merged_ranges
    assert 10 in profile.likely_header_rows
    assert 'B' in profile.numeric_columns
    assert len(profile.blank_rows) <= 100


def saved_template(result):
    return SimpleNamespace(id=1, name='Saved Layout', worksheet=result.worksheet,
        template_type=result.suggested_template_type, header_row=result.table.header_row,
        data_start_row=result.table.data_start_row,
        field_mappings=result.mappings + (result.table.columns if result.suggested_template_type == 'invoice' else []))


def test_template_recognition_ignores_business_values(tmp_path):
    wb = northstar()
    result = analyze(wb, tmp_path)
    template = saved_template(result)
    wb.active['A1'] = 'Completely Different Company'
    wb.active['F5'] = 'OTHER-2027-1234'
    wb.active['E18'] = 999
    assert match_templates(analyze(wb, tmp_path), [template])[0].template_id == 1


def test_template_mismatch_and_generic_invoice_false_positive(tmp_path):
    result = analyze(northstar(), tmp_path)
    template = saved_template(result)
    template.field_mappings = [m.model_copy(update={'cell_ref': 'Z99'}) if m.mapping_type == 'cell' else m for m in template.field_mappings]
    assert match_templates(result, [template]) == []
    template.field_mappings = result.mappings[:1]
    assert match_templates(result, [template]) == []


async def test_multiple_sheets_selection_and_api_errors(client):
    wb = northstar()
    wb.create_sheet('Summary', 0)['A1'] = 'Notes'
    wb.create_sheet('Instructions')['A1'] = 'Please read'
    file_id = await upload(client, wb)
    response = await client.post('/api/v1/suggestions/analyze', json={'file_id': file_id})
    assert response.status_code == 200
    data = response.json()
    assert data['suggested_worksheet'] == 'Invoice'
    assert data['review_required'] is True
    assert len(data['profiles']) == 3
    other = await client.post('/api/v1/suggestions/analyze', json={'file_id': file_id, 'worksheet': 'Summary'})
    assert other.json()['analysis']['confidence'] == 'low'
    assert (await client.post('/api/v1/suggestions/analyze', json={'file_id': file_id, 'worksheet': 'Missing'})).status_code == 404
    assert (await client.post('/api/v1/suggestions/analyze', json={'file_id': 99999})).status_code == 404
    assert (await client.post('/api/v1/suggestions/analyze', json={'file_id': -1})).status_code == 422


@pytest.mark.parametrize('factory', [northstar, dataset])
async def test_suggestions_use_existing_template_and_preview_pipeline(client, factory):
    file_id = await upload(client, factory())
    data = (await client.post('/api/v1/suggestions/analyze', json={'file_id': file_id})).json()
    assert (await client.get('/api/v1/templates')).json() == []  # Analysis never creates templates.
    analysis = data['analysis']
    fields = analysis['mappings'] + (analysis['table']['columns'] if analysis['suggested_template_type'] == 'invoice' else [])
    for field in fields:
        field.pop('confidence')
        field.pop('reason')
    created = await client.post('/api/v1/templates', json={
        'name': 'Reviewed Suggestions', 'template_type': analysis['suggested_template_type'],
        'worksheet': analysis['worksheet'], 'header_row': analysis['table']['header_row'],
        'data_start_row': analysis['table']['data_start_row'], 'field_mappings': fields,
    })
    assert created.status_code == 201, created.text
    template_id = created.json()['id']
    preview = await client.post('/api/v1/imports/extract', json={'file_id': file_id, 'template_id': template_id})
    assert preview.status_code == 200, preview.text
    matches = (await client.post('/api/v1/suggestions/analyze', json={'file_id': file_id})).json()['template_matches']
    assert matches[0]['template_id'] == template_id
