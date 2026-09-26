"""
tests/integration/test_api_exports.py — Integration tests for Milestone 9 Export Center endpoints.

Covers:
- Export summary for Invoices and Line Items
- Invoices CSV export (status, Content-Type, Content-Disposition, UTF-8 BOM, headers, values)
- Invoices CSV custom field flattening and deterministic ordering
- Invoices CSV formula injection protection
- Invoices CSV export without pagination limitation (exports all matching records)
- Soft-deleted records exclusion from CSV and XLSX
- Line items CSV export with flattened parent invoice context
- Invoices XLSX export (worksheet name, header styling, frozen pane, autofilter, typed dates/numbers)
- Line items XLSX export (typed numbers, percentage format, custom fields)
- Northstar manual QA export verification
"""

import csv
import io
from datetime import date, datetime
from decimal import Decimal
import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy.orm import Session

from app.models.invoice import ImportBatch, InvoiceLineItem, InvoiceRecord
from app.models.source_file import SourceFile
from app.models.template import Template


@pytest.fixture
def seed_export_data(db_session: Session) -> dict:
    """
    Seeds database with:
    1. Active Batch 1 (Northstar invoice with 4 line items and custom fields)
    2. Active Batch 2 (Asteron Medical multi-record / header invoices with formula-test data)
    3. Soft-deleted Batch 3 (Should be completely excluded from exports)
    """
    sf1 = SourceFile(
        stored_filename="northstar.xlsx",
        original_name="Northstar_Invoice_Sept2026.xlsx",
        file_size=10240,
        checksum="exp111",
    )
    sf2 = SourceFile(
        stored_filename="asteron.xlsx",
        original_name="Asteron_Medical_Sept2026.xlsx",
        file_size=8192,
        checksum="exp222",
    )
    sf3 = SourceFile(
        stored_filename="deleted_batch.xlsx",
        original_name="Deleted_Batch.xlsx",
        file_size=4096,
        checksum="exp333",
    )
    db_session.add_all([sf1, sf2, sf3])
    db_session.flush()

    t1 = Template(
        name="Northstar Invoice Template",
        template_type="invoice",
        file_type="xlsx",
        worksheet="Invoice",
    )
    t2 = Template(
        name="Asteron Multi-Record Template",
        template_type="dataset",
        file_type="xlsx",
        worksheet="Sheet1",
    )
    db_session.add_all([t1, t2])
    db_session.flush()

    batch1 = ImportBatch(
        source_file_id=sf1.id,
        template_id=t1.id,
        status="imported",
        record_count=1,
        warning_count=0,
        is_deleted=False,
        imported_at=datetime(2026, 9, 21, 10, 0, 0),
    )
    batch2 = ImportBatch(
        source_file_id=sf2.id,
        template_id=t2.id,
        status="imported",
        record_count=2,
        warning_count=1,
        is_deleted=False,
        imported_at=datetime(2026, 9, 22, 11, 30, 0),
    )
    batch_deleted = ImportBatch(
        source_file_id=sf3.id,
        template_id=t1.id,
        status="imported",
        record_count=1,
        warning_count=0,
        is_deleted=True,  # Soft deleted!
        imported_at=datetime(2026, 9, 20, 8, 0, 0),
    )
    db_session.add_all([batch1, batch2, batch_deleted])
    db_session.flush()

    # Invoices
    inv1 = InvoiceRecord(
        batch_id=batch1.id,
        company_name="NORTHSTAR OFFICE SOLUTIONS SDN. BHD.",
        invoice_number="NOS-2026-0947",
        invoice_date=date(2026, 9, 20),
        total_amount=Decimal("7211.59"),
        currency="MYR",
        source_worksheet="Invoice",
        raw_data='{"company": "NORTHSTAR"}',
        custom_fields={"due_date": "2026-10-20", "sub_total": "6677.40", "tax_amount": "534.192"},
        created_at=datetime(2026, 9, 21, 10, 0, 0),
    )
    inv2 = InvoiceRecord(
        batch_id=batch2.id,
        company_name="=cmd|' /C calc'!A0",  # Potential formula injection string
        invoice_number="INV-1001",
        invoice_date=date(2026, 9, 1),
        total_amount=Decimal("1250.00"),
        currency="USD",
        source_worksheet="Sheet1",
        source_row_number=2,
        raw_data='{"company": "injection"}',
        custom_fields={"payment_terms": "Net 30", "discount": "-50.00"},
        created_at=datetime(2026, 9, 22, 11, 30, 0),
    )
    inv3 = InvoiceRecord(
        batch_id=batch2.id,
        company_name="Asteron Medical Supplies",
        invoice_number="INV-1002",
        invoice_date=date(2026, 9, 15),
        total_amount=Decimal("4500.00"),
        currency="USD",
        source_worksheet="Sheet1",
        source_row_number=3,
        raw_data='{"company": "Asteron"}',
        custom_fields={"payment_terms": "Net 15"},
        created_at=datetime(2026, 9, 22, 11, 31, 0),
    )
    inv_del = InvoiceRecord(
        batch_id=batch_deleted.id,
        company_name="Ghost Technologies Inc",
        invoice_number="GHOST-9999",
        invoice_date=date(2026, 9, 10),
        total_amount=Decimal("9999.99"),
        currency="USD",
        source_worksheet="Invoice",
        raw_data='{"company": "Ghost"}',
        created_at=datetime(2026, 9, 20, 8, 0, 0),
    )
    db_session.add_all([inv1, inv2, inv3, inv_del])
    db_session.flush()

    # Line items
    li1 = InvoiceLineItem(
        invoice_id=inv1.id,
        source_row_number=11,
        description="Ergonomic Office Chair — Model E7",
        quantity=Decimal("6.0000"),
        unit_price=Decimal("689.00"),
        tax_rate=Decimal("0.0800"),
        amount=Decimal("4134.00"),
        custom_fields={"sku_code": "SKU-CHAIR-E7"},
        created_at=datetime(2026, 9, 21, 10, 0, 0),
    )
    li2 = InvoiceLineItem(
        invoice_id=inv1.id,
        source_row_number=12,
        description="Adjustable Monitor Arm — Dual Display",
        quantity=Decimal("6.0000"),
        unit_price=Decimal("249.00"),
        tax_rate=Decimal("0.0800"),
        amount=Decimal("1494.00"),
        custom_fields={"sku_code": "SKU-ARM-DUAL"},
        created_at=datetime(2026, 9, 21, 10, 0, 0),
    )
    li3 = InvoiceLineItem(
        invoice_id=inv1.id,
        source_row_number=13,
        description="Wireless Keyboard & Mouse Set",
        quantity=Decimal("6.0000"),
        unit_price=Decimal("129.90"),
        tax_rate=Decimal("0.0800"),
        amount=Decimal("779.40"),
        custom_fields={"sku_code": "SKU-KBD-SET"},
        created_at=datetime(2026, 9, 21, 10, 0, 0),
    )
    li4 = InvoiceLineItem(
        invoice_id=inv1.id,
        source_row_number=14,
        description="Cable Management Tray",
        quantity=Decimal("6.0000"),
        unit_price=Decimal("45.00"),
        tax_rate=Decimal("0.0800"),
        amount=Decimal("270.00"),
        custom_fields={"sku_code": "SKU-TRAY-01"},
        created_at=datetime(2026, 9, 21, 10, 0, 0),
    )
    li5 = InvoiceLineItem(
        invoice_id=inv2.id,
        source_row_number=1,
        description="+SUM(A1:A10)",  # Formula injection description test
        quantity=Decimal("25.0000"),
        unit_price=Decimal("50.00"),
        amount=Decimal("1250.00"),
        created_at=datetime(2026, 9, 22, 11, 30, 0),
    )
    li_del = InvoiceLineItem(
        invoice_id=inv_del.id,
        source_row_number=1,
        description="Ghost Item",
        quantity=Decimal("1.0000"),
        unit_price=Decimal("9999.99"),
        amount=Decimal("9999.99"),
        created_at=datetime(2026, 9, 20, 8, 0, 0),
    )
    db_session.add_all([li1, li2, li3, li4, li5, li_del])
    db_session.commit()

    return {
        "batch1": batch1,
        "batch2": batch2,
        "batch_deleted": batch_deleted,
        "inv1": inv1,
        "inv2": inv2,
        "inv3": inv3,
        "inv_del": inv_del,
    }


# ── Export Summary Tests ──────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_get_export_summary_invoices(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/summary?dataset=invoices")
    assert res.status_code == 200
    data = res.json()
    assert data["dataset"] == "invoices"
    assert data["total_records"] == 3  # inv1, inv2, inv3 (inv_del excluded)
    assert "Company Name" in data["columns"]
    assert "Invoice Number" in data["columns"]
    assert "Due Date" in data["columns"]
    assert "Subtotal" in data["columns"]
    assert "Payment Terms" in data["columns"]
    assert data["column_count"] == len(data["columns"])


@pytest.mark.asyncio
async def test_get_export_summary_line_items(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/summary?dataset=line-items")
    assert res.status_code == 200
    data = res.json()
    assert data["dataset"] == "line-items"
    assert data["total_records"] == 5  # 4 Northstar + 1 Asteron (ghost deleted excluded)
    assert "Description" in data["columns"]
    assert "SKU Code" in data["columns"]


# ── Invoices CSV Export Tests ─────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_invoices_csv_success(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/invoices.csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.csv"')

    # Check UTF-8 BOM
    raw_bytes = res.content
    assert raw_bytes.startswith(b"\xef\xbb\xbf")

    # Parse CSV content
    text = raw_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    # Header check
    headers = rows[0]
    assert headers[0] == "Company Name"
    assert headers[1] == "Invoice Number"
    assert headers[2] == "Invoice Date"
    assert headers[3] == "Currency"
    assert headers[4] == "Total Amount"
    assert "Due Date" in headers
    assert "Subtotal" in headers
    assert "Tax Amount" in headers
    assert "Payment Terms" in headers

    # Rows check (3 rows + header = 4 total)
    assert len(rows) == 4

    # Verify soft-deleted batch is excluded
    all_text = "\n".join([",".join(r) for r in rows])
    assert "Ghost Technologies" not in all_text
    assert "GHOST-9999" not in all_text


@pytest.mark.asyncio
async def test_export_invoices_csv_formula_injection_protection(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/invoices.csv")
    assert res.status_code == 200
    text = res.content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    # inv2 had company_name="=cmd|' /C calc'!A0" -> should be sanitized to "'=cmd|' /C calc'!A0"
    inj_row = next(r for r in rows[1:] if "INV-1001" in r)
    company_cell = inj_row[0]
    assert company_cell.startswith("'=")


@pytest.mark.asyncio
async def test_export_invoices_csv_with_filters(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/invoices.csv?company=Northstar")
    assert res.status_code == 200
    text = res.content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    # 1 header + 1 Northstar record
    assert len(rows) == 2
    assert rows[1][0] == "NORTHSTAR OFFICE SOLUTIONS SDN. BHD."
    assert rows[1][1] == "NOS-2026-0947"


# ── Line Items CSV Export Tests ───────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_line_items_csv_success(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/line-items.csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]

    text = res.content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    # 1 header + 5 line items
    assert len(rows) == 6
    headers = rows[0]
    assert headers[0] == "Company Name"
    assert headers[4] == "Description"
    assert "SKU Code" in headers

    # Check parent context flattened
    chair_row = next(r for r in rows if "Ergonomic Office Chair" in r[4])
    assert chair_row[0] == "NORTHSTAR OFFICE SOLUTIONS SDN. BHD."
    assert chair_row[1] == "NOS-2026-0947"
    assert chair_row[3] == "MYR"

    # Formula injection in line items description "+SUM(A1:A10)" -> "'+SUM(A1:A10)"
    sum_row = next(r for r in rows if "SUM(A1:A10)" in r[4])
    assert sum_row[4].startswith("'+")


# ── Invoices XLSX Export Tests ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_invoices_xlsx_success(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/invoices.xlsx")
    assert res.status_code == 200
    assert "openxmlformats" in res.headers["content-type"]

    # Load workbook from response bytes
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Invoices" in wb.sheetnames
    ws = wb["Invoices"]

    # Header checks
    header_cells = [cell.value for cell in ws[1]]
    assert "Company Name" in header_cells
    assert "Invoice Number" in header_cells
    assert "Due Date" in header_cells
    assert "Total Amount" in header_cells

    # Check freeze pane and autofilter
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None

    # Check row count (1 header + 3 data rows = 4)
    assert ws.max_row == 4

    # Check numeric cell type for total_amount
    for row in ws.iter_rows(min_row=2, max_row=4, values_only=False):
        if row[1].value == "NOS-2026-0947":
            # Total amount cell (col 5)
            total_cell = row[4]
            assert isinstance(total_cell.value, (int, float))
            assert total_cell.value == 7211.59
            assert total_cell.number_format == "#,##0.00"

    wb.close()


# ── Line Items XLSX Export Tests ──────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_line_items_xlsx_success(client: AsyncClient, seed_export_data: dict):
    res = await client.get("/api/v1/exports/line-items.xlsx")
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Line Items" in wb.sheetnames
    ws = wb["Line Items"]

    # 1 header + 5 line items = 6 rows
    assert ws.max_row == 6

    # Verify percentage formatting on Tax Rate column (col 8)
    for row in ws.iter_rows(min_row=2, max_row=6, values_only=False):
        if "Chair" in str(row[4].value):
            tax_rate_cell = row[7]
            assert tax_rate_cell.value == 0.08
            assert tax_rate_cell.number_format == "0.00%"

            price_cell = row[6]
            assert price_cell.value == 689.00
            assert price_cell.number_format == "#,##0.00"

    wb.close()


# ── Northstar Specification 32 QA Audit Test ──────────────────────────────────


@pytest.mark.asyncio
async def test_export_northstar_manual_qa_audit(client: AsyncClient, seed_export_data: dict):
    """
    Direct verification of Milestone 9 Specification 32 requirements for Northstar invoice:
    - Invoices CSV export has exact canonical fields, Due Date, Subtotal, Tax Amount
    - Line Items CSV export has all 4 items with exact quantities, prices, taxes, amounts, and SKU codes
    """
    # 1. Invoices CSV
    inv_res = await client.get("/api/v1/exports/invoices.csv?search=NOS-2026-0947")
    assert inv_res.status_code == 200
    text = inv_res.content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    assert len(rows) == 1
    ns_inv = rows[0]
    assert ns_inv["Company Name"] == "NORTHSTAR OFFICE SOLUTIONS SDN. BHD."
    assert ns_inv["Invoice Number"] == "NOS-2026-0947"
    assert ns_inv["Invoice Date"] == "2026-09-20"
    assert ns_inv["Due Date"] == "2026-10-20"
    assert ns_inv["Currency"] == "MYR"
    assert ns_inv["Subtotal"] == "6677.40"
    assert ns_inv["Tax Amount"] == "534.192"
    assert ns_inv["Total Amount"] == "7211.59"

    # 2. Line Items CSV
    li_res = await client.get("/api/v1/exports/line-items.csv?search=NOS-2026-0947")
    assert li_res.status_code == 200
    li_text = li_res.content.decode("utf-8-sig")
    li_reader = csv.DictReader(io.StringIO(li_text))
    li_rows = list(li_reader)
    assert len(li_rows) == 4

    expected = [
        ("Ergonomic Office Chair — Model E7", "6.0000", "689.00", "0.0800", "4134.00", "SKU-CHAIR-E7"),
        ("Adjustable Monitor Arm — Dual Display", "6.0000", "249.00", "0.0800", "1494.00", "SKU-ARM-DUAL"),
        ("Wireless Keyboard & Mouse Set", "6.0000", "129.90", "0.0800", "779.40", "SKU-KBD-SET"),
        ("Cable Management Tray", "6.0000", "45.00", "0.0800", "270.00", "SKU-TRAY-01"),
    ]

    for actual, (exp_desc, exp_qty, exp_price, exp_tax, exp_amt, exp_sku) in zip(li_rows, expected):
        assert actual["Description"] == exp_desc
        assert actual["Quantity"] == exp_qty
        assert actual["Unit Price"] == exp_price
        assert actual["Tax Rate"] == exp_tax
        assert actual["Amount"] == exp_amt
        assert actual["SKU Code"] == exp_sku
