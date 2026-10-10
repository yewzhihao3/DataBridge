"""
tests/integration/test_m16_3_1_source_preview_hotfix.py — Comprehensive tests for M16.3.1 Source Preview Reliability Hotfix.

Covers:
1. Invoice from XLSX with worksheet name
2. XLSX with multiple worksheets
3. XLSX with spaces/special characters in sheet names
4. Missing source worksheet
5. Invalid worksheet name
6. Receipt JPG source
7. Receipt PNG source
8. PDF invoice source
9. PDF receipt source
10. Missing source file
11. Unauthorized source access
12. Cross-tenant source access
13. Existing Data Explorer navigation
14. Existing export functionality
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.orm import Session

from app.config import settings
from app.main import app
from app.models.identity import Organization, OrganizationMembership, User
from app.models.invoice import ImportBatch, InvoiceLineItem, InvoiceRecord, ReceiptLineItem, ReceiptRecord
from app.models.source_file import SourceFile
from app.models.template import Template
from app.security import password_hasher
from app.services.storage import storage
from tests.conftest import TestingSessionLocal


@pytest.fixture
def make_xlsx():
    """Helper to create an XLSX file with specified sheet names and data."""
    def _create(sheets_data: dict[str, list[list[str]]]) -> bytes:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)
        for sheet_name, rows in sheets_data.items():
            ws = wb.create_sheet(title=sheet_name)
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=val)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    return _create


@pytest.fixture
def seed_preview_test_data(db_session: Session, make_xlsx) -> dict:
    """Seeds multi-type source files and corresponding invoice/receipt business records."""
    # 1. XLSX file with multiple sheets (including spaces & special characters)
    xlsx_bytes = make_xlsx({
        "Invoice": [["Invoice #", "Amount"], ["INV-001", "100.00"]],
        "Sales 2024": [["Date", "Customer", "Total"], ["2024-01-01", "Acme", "500.00"]],
        "Invoice & Receipts": [["Code", "Desc"], ["A1", "Item One"]],
    })
    key_xlsx = storage.new_key(False, ".xlsx")
    with open(storage.local_path(key_xlsx), "wb") as f:
        f.write(xlsx_bytes)

    sf_xlsx = SourceFile(
        organization_id=1,
        original_name="Invoices_MultiSheet.xlsx",
        stored_filename=key_xlsx,
        file_size=len(xlsx_bytes),
        checksum="chk_xlsx_1",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    # 2. Receipt JPG file
    jpg_bytes = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x01\x00`\x00`\x00\x00\xff\xdb"
    key_jpg = storage.new_key(False, ".jpg")
    with open(storage.local_path(key_jpg), "wb") as f:
        f.write(jpg_bytes)

    sf_jpg = SourceFile(
        organization_id=1,
        original_name="Store_Receipt.jpg",
        stored_filename=key_jpg,
        file_size=len(jpg_bytes),
        checksum="chk_jpg_1",
        media_type="image/jpeg",
    )

    # 3. Receipt PNG file
    png_bytes = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4"
    key_png = storage.new_key(False, ".png")
    with open(storage.local_path(key_png), "wb") as f:
        f.write(png_bytes)

    sf_png = SourceFile(
        organization_id=1,
        original_name="Receipt_Scan.png",
        stored_filename=key_png,
        file_size=len(png_bytes),
        checksum="chk_png_1",
        media_type="image/png",
    )

    # 4. PDF invoice file
    pdf_inv_bytes = b"%PDF-1.4\n%test invoice pdf\n%%EOF"
    key_pdf_inv = storage.new_key(False, ".pdf")
    with open(storage.local_path(key_pdf_inv), "wb") as f:
        f.write(pdf_inv_bytes)

    sf_pdf_inv = SourceFile(
        organization_id=1,
        original_name="Vendor_Invoice.pdf",
        stored_filename=key_pdf_inv,
        file_size=len(pdf_inv_bytes),
        checksum="chk_pdf_inv_1",
        media_type="application/pdf",
    )

    # 5. PDF receipt file
    pdf_rec_bytes = b"%PDF-1.4\n%test receipt pdf\n%%EOF"
    key_pdf_rec = storage.new_key(False, ".pdf")
    with open(storage.local_path(key_pdf_rec), "wb") as f:
        f.write(pdf_rec_bytes)

    sf_pdf_rec = SourceFile(
        organization_id=1,
        original_name="Gas_Receipt.pdf",
        stored_filename=key_pdf_rec,
        file_size=len(pdf_rec_bytes),
        checksum="chk_pdf_rec_1",
        media_type="application/pdf",
    )

    db_session.add_all([sf_xlsx, sf_jpg, sf_png, sf_pdf_inv, sf_pdf_rec])
    db_session.flush()

    # Templates
    tmpl_xlsx = Template(
        organization_id=1,
        name="Excel Invoices",
        template_type="invoice",
        file_type="xlsx",
        worksheet="Invoice",
    )
    tmpl_doc = Template(
        organization_id=1,
        name="Document Invoices",
        template_type="invoice",
        file_type="document",
        worksheet="Document",
    )
    db_session.add_all([tmpl_xlsx, tmpl_doc])
    db_session.flush()

    # Batch 1 (Excel Invoices)
    batch_xlsx = ImportBatch(
        organization_id=1,
        source_file_id=sf_xlsx.id,
        template_id=tmpl_xlsx.id,
        status="imported",
        record_count=2,
        is_deleted=False,
    )
    # Batch 2 (JPG Receipt)
    batch_jpg = ImportBatch(
        organization_id=1,
        source_file_id=sf_jpg.id,
        template_id=tmpl_doc.id,
        status="imported",
        record_count=1,
        is_deleted=False,
    )
    # Batch 3 (PNG Receipt)
    batch_png = ImportBatch(
        organization_id=1,
        source_file_id=sf_png.id,
        template_id=tmpl_doc.id,
        status="imported",
        record_count=1,
        is_deleted=False,
    )
    # Batch 4 (PDF Invoice)
    batch_pdf_inv = ImportBatch(
        organization_id=1,
        source_file_id=sf_pdf_inv.id,
        template_id=tmpl_doc.id,
        status="imported",
        record_count=1,
        is_deleted=False,
    )
    # Batch 5 (PDF Receipt)
    batch_pdf_rec = ImportBatch(
        organization_id=1,
        source_file_id=sf_pdf_rec.id,
        template_id=tmpl_doc.id,
        status="imported",
        record_count=1,
        is_deleted=False,
    )

    db_session.add_all([batch_xlsx, batch_jpg, batch_png, batch_pdf_inv, batch_pdf_rec])
    db_session.flush()

    # Invoices
    inv_xlsx_1 = InvoiceRecord(
        batch_id=batch_xlsx.id,
        company_name="Apex Global",
        invoice_number="INV-APX-01",
        invoice_date=date(2026, 9, 20),
        total_amount=Decimal("1200.00"),
        currency="USD",
        direction="sales",
        source_worksheet="Invoice",
        raw_data="{}",
    )
    inv_xlsx_2 = InvoiceRecord(
        batch_id=batch_xlsx.id,
        company_name="Specialty Co",
        invoice_number="INV-SPC-02",
        invoice_date=date(2026, 9, 21),
        total_amount=Decimal("450.00"),
        currency="USD",
        direction="sales",
        source_worksheet="Sales 2024",
        raw_data="{}",
    )
    inv_pdf = InvoiceRecord(
        batch_id=batch_pdf_inv.id,
        company_name="PDF Supplier Inc",
        invoice_number="INV-PDF-99",
        invoice_date=date(2026, 9, 22),
        total_amount=Decimal("890.00"),
        currency="USD",
        direction="purchase",
        source_worksheet="Document",
        raw_data="{}",
    )

    # Receipts
    rec_jpg = ReceiptRecord(
        batch_id=batch_jpg.id,
        merchant_name="QuickMart",
        receipt_number="REC-QM-1",
        transaction_date=date(2026, 9, 20),
        total_amount=Decimal("45.50"),
        currency="USD",
        direction="expense",
        payment_method="Card",
        raw_data="{}",
    )
    rec_png = ReceiptRecord(
        batch_id=batch_png.id,
        merchant_name="Office World",
        receipt_number="REC-OW-2",
        transaction_date=date(2026, 9, 21),
        total_amount=Decimal("120.00"),
        currency="USD",
        direction="expense",
        payment_method="Cash",
        raw_data="{}",
    )
    rec_pdf = ReceiptRecord(
        batch_id=batch_pdf_rec.id,
        merchant_name="Shell Petroleum",
        receipt_number="REC-SHL-3",
        transaction_date=date(2026, 9, 22),
        total_amount=Decimal("75.00"),
        currency="USD",
        direction="expense",
        payment_method="Card",
        raw_data="{}",
    )

    db_session.add_all([inv_xlsx_1, inv_xlsx_2, inv_pdf, rec_jpg, rec_png, rec_pdf])
    db_session.commit()

    return {
        "sf_xlsx": sf_xlsx,
        "sf_jpg": sf_jpg,
        "sf_png": sf_png,
        "sf_pdf_inv": sf_pdf_inv,
        "sf_pdf_rec": sf_pdf_rec,
        "inv_xlsx_1": inv_xlsx_1,
        "inv_xlsx_2": inv_xlsx_2,
        "inv_pdf": inv_pdf,
        "rec_jpg": rec_jpg,
        "rec_png": rec_png,
        "rec_pdf": rec_pdf,
    }


# ── 1. Invoice from XLSX with Worksheet Name ──────────────────────────────────

@pytest.mark.asyncio
async def test_preview_xlsx_invoice_with_worksheet(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]
    inv = seed_preview_test_data["inv_xlsx_1"]

    res = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name={inv.source_worksheet}")
    assert res.status_code == 200
    data = res.json()
    assert data["file_id"] == sf_xlsx.id
    assert data["sheet_name"] == "Invoice"
    assert len(data["rows"]) >= 2
    assert data["rows"][0][0] == "Invoice #"
    assert data["total_rows_returned"] >= 2


# ── 2. XLSX with Multiple Worksheets ──────────────────────────────────────────

@pytest.mark.asyncio
async def test_xlsx_multiple_worksheets_inspection_and_preview(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]

    # Inspection details return all sheets
    inspect_res = await client.get(f"/api/v1/files/{sf_xlsx.id}")
    assert inspect_res.status_code == 200
    sheet_names = [w["name"] for w in inspect_res.json()["inspection"]["worksheets"]]
    assert "Invoice" in sheet_names
    assert "Sales 2024" in sheet_names
    assert "Invoice & Receipts" in sheet_names

    # Preview different sheets individually
    res_sheet1 = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=Invoice")
    assert res_sheet1.status_code == 200
    assert res_sheet1.json()["sheet_name"] == "Invoice"

    res_sheet2 = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=Sales%202024")
    assert res_sheet2.status_code == 200
    assert res_sheet2.json()["sheet_name"] == "Sales 2024"


# ── 3. XLSX with Spaces and Special Characters in Sheet Names ─────────────────

@pytest.mark.asyncio
async def test_xlsx_special_character_sheet_names(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]

    # Sheet with space: "Sales 2024"
    res1 = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=Sales%202024")
    assert res1.status_code == 200
    assert res1.json()["sheet_name"] == "Sales 2024"

    # Sheet with & ampersand: "Invoice & Receipts"
    res2 = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=Invoice%20%26%20Receipts")
    assert res2.status_code == 200
    assert res2.json()["sheet_name"] == "Invoice & Receipts"


# ── 4. Missing Source Worksheet ───────────────────────────────────────────────

@pytest.mark.asyncio
async def test_missing_worksheet_returns_user_friendly_400(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]

    # Calling without sheet_name
    res = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview")
    assert res.status_code == 400
    assert "Worksheet name is required" in res.json()["detail"]

    # Calling with empty sheet_name
    res_empty = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=")
    assert res_empty.status_code == 400
    assert "Worksheet name is required" in res_empty.json()["detail"]


# ── 5. Invalid Worksheet Name ─────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_invalid_worksheet_name_returns_404(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]

    res = await client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=NonExistentSheet")
    assert res.status_code == 404
    assert "not found" in res.json()["detail"].lower()


# ── 6. Receipt JPG Source ─────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_receipt_jpg_preview(client: AsyncClient, seed_preview_test_data):
    sf_jpg = seed_preview_test_data["sf_jpg"]

    res = await client.get(f"/api/v1/documents/{sf_jpg.id}/preview")
    assert res.status_code == 200
    assert "image/jpeg" in res.headers["content-type"]
    assert res.headers.get("content-disposition", "").startswith("inline")


# ── 7. Receipt PNG Source ─────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_receipt_png_preview(client: AsyncClient, seed_preview_test_data):
    sf_png = seed_preview_test_data["sf_png"]

    res = await client.get(f"/api/v1/documents/{sf_png.id}/preview")
    assert res.status_code == 200
    assert "image/png" in res.headers["content-type"]
    assert res.headers.get("content-disposition", "").startswith("inline")


# ── 8. PDF Invoice Source ─────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_pdf_invoice_preview(client: AsyncClient, seed_preview_test_data):
    sf_pdf_inv = seed_preview_test_data["sf_pdf_inv"]

    res = await client.get(f"/api/v1/documents/{sf_pdf_inv.id}/preview")
    assert res.status_code == 200
    assert "application/pdf" in res.headers["content-type"]
    assert res.headers.get("content-disposition", "").startswith("inline")


# ── 9. PDF Receipt Source ─────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_pdf_receipt_preview(client: AsyncClient, seed_preview_test_data):
    sf_pdf_rec = seed_preview_test_data["sf_pdf_rec"]

    res = await client.get(f"/api/v1/documents/{sf_pdf_rec.id}/preview")
    assert res.status_code == 200
    assert "application/pdf" in res.headers["content-type"]
    assert res.headers.get("content-disposition", "").startswith("inline")


# ── 10. Missing Source File ───────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_missing_source_file_returns_404(client: AsyncClient):
    # Preview missing file
    res1 = await client.get("/api/v1/files/99999/preview?sheet_name=Sheet1")
    assert res1.status_code == 404

    # Download missing file
    res2 = await client.get("/api/v1/files/99999/download")
    assert res2.status_code == 404

    # Preview missing document
    res3 = await client.get("/api/v1/documents/99999/preview")
    assert res3.status_code == 404

    # Download missing document
    res4 = await client.get("/api/v1/documents/99999/download")
    assert res4.status_code == 404


# ── 11. Unauthorized Source Access ────────────────────────────────────────────

@pytest.mark.asyncio
async def test_unauthorized_source_access(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]
    sf_jpg = seed_preview_test_data["sf_jpg"]

    # Anonymous client with no auth cookie
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as anon_client:
        res1 = await anon_client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=Invoice")
        assert res1.status_code == 401

        res2 = await anon_client.get(f"/api/v1/files/{sf_xlsx.id}/download")
        assert res2.status_code == 401

        res3 = await anon_client.get(f"/api/v1/documents/{sf_jpg.id}/preview")
        assert res3.status_code == 401

        res4 = await anon_client.get(f"/api/v1/documents/{sf_jpg.id}/download")
        assert res4.status_code == 401


# ── 12. Cross-Tenant Source Access ────────────────────────────────────────────

@pytest.mark.asyncio
async def test_cross_tenant_source_access_blocked(client: AsyncClient, seed_preview_test_data):
    sf_xlsx = seed_preview_test_data["sf_xlsx"]
    sf_jpg = seed_preview_test_data["sf_jpg"]

    # Create Organization 2 and User 2
    with TestingSessionLocal() as session:
        org2 = Organization(id=2, name="Other Tenant", slug="other-tenant")
        session.add(org2)
        session.flush()

        user2 = User(
            email="tenant2@example.test",
            display_name="Tenant 2 User",
            password_hash=password_hasher.hash("tenant2-password"),
        )
        session.add(user2)
        session.flush()

        session.add(OrganizationMembership(user_id=user2.id, organization_id=2, role="OWNER"))
        session.commit()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as t2_client:
        login = await t2_client.post(
            "/api/v1/auth/login",
            json={"email": "tenant2@example.test", "password": "tenant2-password"},
        )
        assert login.status_code == 200
        t2_client.headers["x-csrf-token"] = login.json()["csrf_token"]
        t2_client.headers["x-workspace-id"] = "2"

        # Tenant 2 attempts to preview Org 1's Excel file
        res1 = await t2_client.get(f"/api/v1/files/{sf_xlsx.id}/preview?sheet_name=Invoice")
        assert res1.status_code == 404

        # Tenant 2 attempts to download Org 1's Excel file
        res2 = await t2_client.get(f"/api/v1/files/{sf_xlsx.id}/download")
        assert res2.status_code == 404

        # Tenant 2 attempts to preview Org 1's Document
        res3 = await t2_client.get(f"/api/v1/documents/{sf_jpg.id}/preview")
        assert res3.status_code == 404

        # Tenant 2 attempts to download Org 1's Document
        res4 = await t2_client.get(f"/api/v1/documents/{sf_jpg.id}/download")
        assert res4.status_code == 404


# ── 13. Existing Data Explorer Navigation ─────────────────────────────────────

@pytest.mark.asyncio
async def test_existing_data_explorer_navigation_and_metadata(client: AsyncClient, seed_preview_test_data):
    # Invoices listing includes source_worksheet and source_media_type
    res_inv = await client.get("/api/v1/data/invoices")
    assert res_inv.status_code == 200
    inv_data = res_inv.json()
    assert inv_data["total"] >= 3
    first_inv = inv_data["items"][0]
    assert "source_worksheet" in first_inv
    assert "source_media_type" in first_inv

    # Receipts listing includes source_media_type
    res_rec = await client.get("/api/v1/data/receipts")
    assert res_rec.status_code == 200
    rec_data = res_rec.json()
    assert rec_data["total"] >= 3
    first_rec = rec_data["items"][0]
    assert "source_media_type" in first_rec

    # Unified documents listing includes source_worksheet and source_media_type
    res_doc = await client.get("/api/v1/data/documents")
    assert res_doc.status_code == 200
    doc_data = res_doc.json()
    assert doc_data["total"] >= 6
    for item in doc_data["items"]:
        assert "source_worksheet" in item
        assert "source_media_type" in item


# ── 14. Existing Export Functionality ─────────────────────────────────────────

@pytest.mark.asyncio
async def test_existing_export_functionality(client: AsyncClient, seed_preview_test_data):
    # Invoices CSV export
    res_inv_csv = await client.get("/api/v1/exports/invoices.csv")
    assert res_inv_csv.status_code == 200
    assert "text/csv" in res_inv_csv.headers["content-type"]
    assert "Apex Global" in res_inv_csv.text

    # Invoices XLSX export
    res_inv_xlsx = await client.get("/api/v1/exports/invoices.xlsx")
    assert res_inv_xlsx.status_code == 200
    assert "openxmlformats" in res_inv_xlsx.headers["content-type"]

    # Receipts CSV export
    res_rec_csv = await client.get("/api/v1/exports/receipts.csv")
    assert res_rec_csv.status_code == 200
    assert "text/csv" in res_rec_csv.headers["content-type"]
    assert "QuickMart" in res_rec_csv.text

    # Documents CSV export
    res_doc_csv = await client.get("/api/v1/exports/documents.csv")
    assert res_doc_csv.status_code == 200
    assert "text/csv" in res_doc_csv.headers["content-type"]
