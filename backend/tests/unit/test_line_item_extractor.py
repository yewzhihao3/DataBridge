"""
tests/unit/test_line_item_extractor.py — Unit tests for composite invoice line item extraction.
"""

from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
import openpyxl
import pytest

from app.services.extractor import (
    extract_from_file,
    extract_from_workbook,
    ExtractionResult,
)


@pytest.fixture
def composite_invoice_file(tmp_path: Path) -> Path:
    """
    Creates a realistic composite invoice Excel workbook:
    - Header in cells (A1: Supplier, F5: Invoice No, F6: Date, F8: Currency, E18: Total)
    - Line items in rows 11-14 (Col A: Description, Col B: Qty, Col C: Unit Price, Col D: Tax Rate, Col E: Amount, Col F: SKU)
    - Row 15: Blank
    - Row 16: Subtotal
    - Row 17: Tax
    - Row 18: Total
    """
    file_path = tmp_path / "composite_invoice.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"

    # Header cells
    ws["A1"] = "NORTHSTAR OFFICE SOLUTIONS"
    ws["A6"] = "Customer Corp"
    ws["F5"] = "NOS-2026-0947"
    ws["F6"] = "2026-09-20"
    ws["F8"] = "MYR"

    # Line items table header at Row 10
    ws["A10"] = "Description"
    ws["B10"] = "Qty"
    ws["C10"] = "Unit Price"
    ws["D10"] = "Tax Rate"
    ws["E10"] = "Amount"
    ws["F10"] = "SKU"

    # Row 11: Product 1
    ws["A11"] = "Ergonomic Office Chair"
    ws["B11"] = 6
    ws["C11"] = 689.00
    ws["D11"] = "8%"
    ws["E11"] = 4134.00
    ws["F11"] = "SKU-CHAIR-01"

    # Row 12: Product 2 (Total Care Cleaning Kit - test that 'Total' inside product name does not terminate!)
    ws["A12"] = "Total Care Cleaning Kit"
    ws["B12"] = 6
    ws["C12"] = 249.00
    ws["D12"] = "8%"
    ws["E12"] = 1494.00
    ws["F12"] = "SKU-CLEAN-02"

    # Row 13: Product 3 (Tax Preparation Service - test that 'Tax' inside product name does not terminate!)
    ws["A13"] = "Tax Preparation Service"
    ws["B13"] = 6
    ws["C13"] = 129.90
    ws["D13"] = "8%"
    ws["E13"] = 779.40
    ws["F13"] = "SKU-SERV-03"

    # Row 14: Product 4
    ws["A14"] = "Cable Management Tray"
    ws["B14"] = 6
    ws["C14"] = 45.00
    ws["D14"] = "8%"
    ws["E14"] = 270.00
    ws["F14"] = "SKU-TRAY-04"

    # Row 15: Blank row

    # Row 16: Summary Subtotal (Footer)
    ws["A16"] = "Subtotal"
    ws["E16"] = 6677.40

    # Row 17: Summary Tax (Footer)
    ws["A17"] = "Tax"
    ws["E17"] = 534.19

    # Row 18: Summary Total (Footer)
    ws["A18"] = "Total"
    ws["E18"] = 7211.59

    wb.save(file_path)
    wb.close()
    return file_path


def test_composite_invoice_extraction_with_line_items(composite_invoice_file: Path):
    template = SimpleNamespace(
        name="Composite Invoice Template",
        template_type="invoice",
        worksheet="Invoice",
        header_row=10,
        data_start_row=11,
        date_format=None,
        field_mappings=[
            # Header mappings
            SimpleNamespace(field_name="supplier", target_field="company_name", mapping_group="header", mapping_type="cell", cell_ref="A1", is_required=True, data_type="text"),
            SimpleNamespace(field_name="inv_no", target_field="invoice_number", mapping_group="header", mapping_type="cell", cell_ref="F5", is_required=True, data_type="text"),
            SimpleNamespace(field_name="inv_date", target_field="invoice_date", mapping_group="header", mapping_type="cell", cell_ref="F6", is_required=True, data_type="date"),
            SimpleNamespace(field_name="currency", target_field="currency", mapping_group="header", mapping_type="cell", cell_ref="F8", is_required=False, data_type="text"),
            SimpleNamespace(field_name="total_amount", target_field="total_amount", mapping_group="header", mapping_type="cell", cell_ref="E18", is_required=True, data_type="decimal"),
            # Line item mappings
            SimpleNamespace(field_name="desc", target_field="description", mapping_group="line_item", mapping_type="column", column_ref="A", is_required=True, data_type="text"),
            SimpleNamespace(field_name="qty", target_field="quantity", mapping_group="line_item", mapping_type="column", column_ref="B", is_required=True, data_type="decimal"),
            SimpleNamespace(field_name="price", target_field="unit_price", mapping_group="line_item", mapping_type="column", column_ref="C", is_required=True, data_type="decimal"),
            SimpleNamespace(field_name="tax_rate", target_field="tax_rate", mapping_group="line_item", mapping_type="column", column_ref="D", is_required=False, data_type="decimal"),
            SimpleNamespace(field_name="amount", target_field="amount", mapping_group="line_item", mapping_type="column", column_ref="E", is_required=True, data_type="decimal"),
            SimpleNamespace(field_name="sku", target_field="sku", mapping_group="line_item", mapping_type="column", column_ref="F", is_required=False, data_type="text"),
        ],
    )

    res = extract_from_file(composite_invoice_file, template)

    assert not res.has_errors
    assert res.error_count == 0
    assert not res.is_multi_record
    assert res.has_line_items is True

    # Check header fields
    header_dict = {f.field_name: f.normalized_value for f in res.fields}
    assert header_dict["supplier"] == "NORTHSTAR OFFICE SOLUTIONS"
    assert header_dict["inv_no"] == "NOS-2026-0947"
    assert header_dict["currency"] == "MYR"
    assert header_dict["total_amount"] == Decimal("7211.59")

    # Check line items: exactly 4 items extracted (Subtotal/Tax/Total footer rows NOT extracted)
    assert len(res.line_items) == 4

    # Item 1
    item1_fields = {f.field_name: f.normalized_value for f in res.line_items[0].fields}
    assert res.line_items[0].row_number == 11
    assert item1_fields["desc"] == "Ergonomic Office Chair"
    assert item1_fields["qty"] == Decimal("6")
    assert item1_fields["price"] == Decimal("689.00")
    assert item1_fields["tax_rate"] == Decimal("8")
    assert item1_fields["amount"] == Decimal("4134.00")
    assert item1_fields["sku"] == "SKU-CHAIR-01"

    # Item 2: "Total Care Cleaning Kit"
    item2_fields = {f.field_name: f.normalized_value for f in res.line_items[1].fields}
    assert res.line_items[1].row_number == 12
    assert item2_fields["desc"] == "Total Care Cleaning Kit"
    assert item2_fields["qty"] == Decimal("6")
    assert item2_fields["price"] == Decimal("249.00")

    # Item 3: "Tax Preparation Service"
    item3_fields = {f.field_name: f.normalized_value for f in res.line_items[2].fields}
    assert res.line_items[2].row_number == 13
    assert item3_fields["desc"] == "Tax Preparation Service"
    assert item3_fields["qty"] == Decimal("6")
    assert item3_fields["price"] == Decimal("129.90")

    # Item 4
    item4_fields = {f.field_name: f.normalized_value for f in res.line_items[3].fields}
    assert res.line_items[3].row_number == 14
    assert item4_fields["desc"] == "Cable Management Tray"
    assert item4_fields["qty"] == Decimal("6")
    assert item4_fields["price"] == Decimal("45.00")


def test_line_item_provenance_coordinates(composite_invoice_file: Path):
    template = SimpleNamespace(
        name="Provenance Test",
        template_type="invoice",
        worksheet="Invoice",
        header_row=10,
        data_start_row=11,
        date_format=None,
        field_mappings=[
            SimpleNamespace(field_name="supplier", target_field="company_name", mapping_group="header", mapping_type="cell", cell_ref="A1", is_required=True, data_type="text"),
            SimpleNamespace(field_name="inv_no", target_field="invoice_number", mapping_group="header", mapping_type="cell", cell_ref="F5", is_required=True, data_type="text"),
            SimpleNamespace(field_name="desc", target_field="description", mapping_group="line_item", mapping_type="column", column_ref="A", is_required=True, data_type="text"),
            SimpleNamespace(field_name="price", target_field="unit_price", mapping_group="line_item", mapping_type="column", column_ref="C", is_required=True, data_type="decimal"),
        ],
    )

    res = extract_from_file(composite_invoice_file, template)
    assert len(res.line_items) == 4

    row1_fields = {f.field_name: f for f in res.line_items[0].fields}
    assert row1_fields["desc"].source_cell_ref == "A11"
    assert row1_fields["price"].source_cell_ref == "C11"

    row4_fields = {f.field_name: f for f in res.line_items[3].fields}
    assert row4_fields["desc"].source_cell_ref == "A14"
    assert row4_fields["price"].source_cell_ref == "C14"
