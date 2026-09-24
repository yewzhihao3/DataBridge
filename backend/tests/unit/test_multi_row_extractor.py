"""
tests/unit/test_multi_row_extractor.py — Unit tests for multi-row column extraction.
"""

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
import openpyxl
import pytest

from app.services.extractor import extract_from_file, extract_from_workbook


@dataclass
class DummyMapping:
    field_name: str
    target_field: str | None = None
    mapping_type: str = "column"
    cell_ref: str | None = None
    column_ref: str | None = None
    is_required: bool = False
    data_type: str = "text"


@dataclass
class DummyTemplate:
    name: str = "Multi-Row Template"
    worksheet: str | None = "Sheet1"
    header_row: int | None = 1
    data_start_row: int | None = 2
    date_format: str | None = None
    field_mappings: list[DummyMapping] | None = None


def test_multi_row_column_extraction(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Header
    ws["A1"] = "Supplier"
    ws["B1"] = "Invoice No"
    ws["C1"] = "Invoice Date"
    ws["D1"] = "Amount"

    # Row 2
    ws["A2"] = "ABC Sdn Bhd"
    ws["B2"] = "INV-001"
    ws["C2"] = "2026-09-20"
    ws["D2"] = 1200.50

    # Row 3
    ws["A3"] = "XYZ Sdn Bhd"
    ws["B3"] = "INV-002"
    ws["C3"] = "2026-09-21"
    ws["D3"] = 850.00

    file_path = tmp_path / "test_multi.xlsx"
    wb.save(file_path)

    template = DummyTemplate(
        worksheet="Sheet1",
        header_row=1,
        data_start_row=2,
        field_mappings=[
            DummyMapping("supplier", target_field="company_name", mapping_type="column", column_ref="A", is_required=True),
            DummyMapping("inv_no", target_field="invoice_number", mapping_type="column", column_ref="B", is_required=True),
            DummyMapping("inv_date", target_field="invoice_date", mapping_type="column", column_ref="C", data_type="date"),
            DummyMapping("amount", target_field="total_amount", mapping_type="column", column_ref="D", data_type="decimal"),
        ],
    )

    result = extract_from_file(file_path, template)

    assert result.is_multi_record is True
    assert result.has_errors is False
    assert len(result.rows) == 2

    row1 = result.rows[0]
    assert row1.row_number == 2
    f1 = {f.field_name: f for f in row1.fields}
    assert f1["supplier"].normalized_value == "ABC Sdn Bhd"
    assert f1["inv_no"].normalized_value == "INV-001"
    assert f1["inv_date"].normalized_value == date(2026, 9, 20)
    assert f1["amount"].normalized_value == Decimal("1200.50")

    row2 = result.rows[1]
    assert row2.row_number == 3
    f2 = {f.field_name: f for f in row2.fields}
    assert f2["supplier"].normalized_value == "XYZ Sdn Bhd"
    assert f2["inv_no"].normalized_value == "INV-002"


def test_multi_row_skips_blank_rows_and_handles_partial_rows(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Row 2: Valid
    ws["A2"] = "Company A"
    ws["B2"] = "INV-100"

    # Row 3: Completely blank
    # Row 4: Completely blank

    # Row 5: Partial row (missing required invoice number)
    ws["A5"] = "Company B"

    file_path = tmp_path / "test_blank_partial.xlsx"
    wb.save(file_path)

    template = DummyTemplate(
        worksheet="Sheet1",
        header_row=1,
        data_start_row=2,
        field_mappings=[
            DummyMapping("company_name", target_field="company_name", mapping_type="column", column_ref="A", is_required=True),
            DummyMapping("invoice_number", target_field="invoice_number", mapping_type="column", column_ref="B", is_required=True),
        ],
    )

    result = extract_from_file(file_path, template)

    assert result.is_multi_record is True
    # Rows extracted: Row 2 and Row 5 (Row 3 & 4 skipped as completely blank)
    row_numbers = [r.row_number for r in result.rows]
    assert row_numbers == [2, 5]

    row5 = result.rows[1]
    assert row5.row_number == 5
    assert row5.has_errors is True
    assert len(row5.errors) == 1
    assert row5.errors[0].error_type == "required_empty"
